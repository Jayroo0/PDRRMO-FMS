import csv
import io
from datetime import date, datetime, time, timedelta
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.auth import logout, authenticate, login
from django.contrib.admin.models import LogEntry, CHANGE, ADDITION 
from django.contrib.contenttypes.models import ContentType          
from django.db import transaction
from django.db.models import Case, When, Value, IntegerField, Q
from django.core.paginator import Paginator
from django.http import HttpResponse, JsonResponse
from django.template.loader import render_to_string
from django.utils.dateparse import parse_datetime
from django.utils import timezone
from .forms import (
    LAND_MAINTENANCE_CHECKLIST,
    SEACRAFT_MAINTENANCE_CHECKLIST,
    MaintenanceChecklistForm,
    MaintenanceFaultForm,
    FleetIncidentForm,
    FleetAssetRegistrationForm,
    OperatorDetailsForm,
    SeacraftRegistrationForm,
    ScheduledMaintenanceForm,
)
from .models import FleetIncident, Vehicle, VehicleType, Driver, VehicleAsset, OperatorProfile

# =========================================================================
# SYSTEM SECURITY & AUDIT LOG HELPERS
# =========================================================================

def login_view(request):
    if request.user.is_authenticated:
        return redirect('dashboard_portal:dashboard_portal')

    if request.method == 'POST':
        username = request.POST.get('username')
        password = request.POST.get('password')

        if not username or not password:
            messages.error(request, "ACCESS DENIED: Username and password are required.")
            return render(request, 'fleet/login.html')

        user = authenticate(request, username=username, password=password)

        if user is not None:
            login(request, user)
            return redirect('dashboard_portal:dashboard_portal')

        messages.error(request, "ACCESS DENIED: Invalid Username or Password.")

    return render(request, 'fleet/login.html')


def custom_user_logout(request):
    if request.method == 'POST':
        logout(request)
        messages.success(request, "You have been logged out of the PDRRMO operations network.")
    return redirect('dashboard_portal:login')

def check_user_role(user, group_name, keywords):
    if user.is_superuser:
        return True
    
    user_group_names = list(user.groups.values_list('name', flat=True))
    if group_name in user_group_names:
        return True
        
    username_lower = user.username.lower()
    if any(kw in username_lower for kw in keywords):
        return True
        
    return False

def log_action_to_admin(request, object_instance, action_flag, change_message):
    LogEntry.objects.create(
        user_id=request.user.id,
        content_type_id=ContentType.objects.get_for_model(object_instance).id,
        object_id=object_instance.id,
        object_repr=str(object_instance),
        action_flag=action_flag,
        change_message=change_message
    )


def add_disposal_reasons(vehicles):
    pending_vehicle_ids = set(
        vehicles.filter(status='PENDING_DISPOSAL').values_list('id', flat=True)
    )
    disposal_reasons = {}
    if pending_vehicle_ids:
        vehicle_content_type = ContentType.objects.get_for_model(Vehicle)
        reason_entries = LogEntry.objects.filter(
            content_type=vehicle_content_type,
            object_id__in=pending_vehicle_ids,
            change_message__contains='Remarks:',
        ).order_by('-action_time').values_list('object_id', 'change_message')

        for vehicle_id, change_message in reason_entries:
            disposal_reasons.setdefault(
                int(vehicle_id),
                change_message.partition('Remarks:')[2].strip(),
            )

    for vehicle in vehicles:
        if vehicle.id in pending_vehicle_ids:
            vehicle.disposal_reason = disposal_reasons.get(vehicle.id, '')
    return vehicles


def seacraft_vehicle_type_filter():
    return (
        Q(vehicle_type__name__iexact='MARINE')
        | Q(vehicle_type__name__iexact='MARITIME')
    )


def maintenance_checklist_for_vehicle(vehicle):
    if vehicle.vehicle_type and vehicle.vehicle_type.name.upper() in {'MARINE', 'MARITIME'}:
        return SEACRAFT_MAINTENANCE_CHECKLIST
    return LAND_MAINTENANCE_CHECKLIST


def move_overdue_vehicles_to_maintenance(request):
    due_vehicles = Vehicle.objects.filter(
        scheduled_maintenance_date__lte=timezone.localdate(),
        status__in=['OPERATIONAL', 'DEPLOYED'],
    ).select_related('vehicle_type', 'assigned_driver')

    for vehicle in due_vehicles:
        due_type = vehicle.get_scheduled_maintenance_type_display()
        details = (
            f" {vehicle.scheduled_maintenance_description}"
            if vehicle.scheduled_maintenance_description else ''
        )
        vehicle.status = 'MAINTENANCE'
        vehicle.maintenance_problem = f"Scheduled {due_type} is due.{details}"
        vehicle.assigned_driver = None
        vehicle.deployment_location = None
        vehicle.deployment_purpose = None
        vehicle.deployment_time = None
        vehicle.scheduled_maintenance_type = ''
        vehicle.scheduled_maintenance_date = None
        vehicle.scheduled_maintenance_description = ''
        vehicle.save(update_fields=[
            'status',
            'maintenance_problem',
            'assigned_driver',
            'deployment_location',
            'deployment_purpose',
            'deployment_time',
            'scheduled_maintenance_type',
            'scheduled_maintenance_date',
            'scheduled_maintenance_description',
        ])
        log_action_to_admin(
            request,
            vehicle,
            CHANGE,
            f"Scheduled maintenance due: {due_type}{details}. Vehicle moved to maintenance.",
        )
        messages.warning(
            request,
            f"{vehicle.model_name} is now in maintenance because scheduled {due_type.lower()} is due.",
        )


def handle_mechanic_maintenance_action(request, vehicle):
    action_type = request.POST.get('action_type')

    if action_type == 'REPORT_FAULT':
        if vehicle.status not in {'OPERATIONAL', 'DEPLOYED'}:
            messages.error(request, f"{vehicle.model_name} must be operational or deployed before a new maintenance fault can be reported.")
            return

        form = MaintenanceFaultForm(
            request.POST,
            prefix=f'fault-{vehicle.pk}',
        )
        if not form.is_valid():
            for errors in form.errors.values():
                for error in errors:
                    messages.error(request, error)
            return

        vehicle.status = 'MAINTENANCE'
        vehicle.maintenance_problem = (
            f"{form.cleaned_data['fault_type'].replace('_', ' ').title()}: "
            f"{form.cleaned_data['fault_description']}"
        )
        vehicle.assigned_driver = None
        vehicle.deployment_location = None
        vehicle.deployment_purpose = None
        vehicle.deployment_time = None
        vehicle.save()
        log_action_to_admin(
            request,
            vehicle,
            CHANGE,
            f"Reported maintenance fault for {vehicle.model_name}: {vehicle.maintenance_problem}",
        )
        messages.success(request, f"Fault recorded. {vehicle.model_name} is now in maintenance.")
        return

    if action_type == 'COMPLETE_MAINTENANCE':
        if vehicle.status != 'MAINTENANCE':
            messages.error(request, f"{vehicle.model_name} must be in maintenance before it can be returned to service.")
            return

        form = MaintenanceChecklistForm(
            request.POST,
            checklist=maintenance_checklist_for_vehicle(vehicle),
            prefix=f'checklist-{vehicle.pk}',
        )
        if not form.is_valid():
            for errors in form.errors.values():
                for error in errors:
                    messages.error(request, error)
            return

        vehicle.status = 'OPERATIONAL'
        vehicle.maintenance_problem = ''
        vehicle.save()
        log_action_to_admin(
            request,
            vehicle,
            CHANGE,
            f"Completed maintenance checklist and returned {vehicle.model_name} to operational status.",
        )
        messages.success(request, f"Checklist complete. {vehicle.model_name} is operational.")
        return

    if action_type == 'SCHEDULE_MAINTENANCE':
        if vehicle.status != 'OPERATIONAL':
            messages.error(request, f"{vehicle.model_name} must be operational before maintenance can be scheduled.")
            return

        form = ScheduledMaintenanceForm(
            request.POST,
            prefix=f'schedule-{vehicle.pk}',
        )
        if not form.is_valid():
            for errors in form.errors.values():
                for error in errors:
                    messages.error(request, error)
            return

        vehicle.scheduled_maintenance_type = form.cleaned_data['maintenance_type']
        vehicle.scheduled_maintenance_date = form.cleaned_data['due_date']
        vehicle.scheduled_maintenance_description = form.cleaned_data['description']
        vehicle.save(update_fields=[
            'scheduled_maintenance_type',
            'scheduled_maintenance_date',
            'scheduled_maintenance_description',
        ])
        log_action_to_admin(
            request,
            vehicle,
            CHANGE,
            f"Scheduled {vehicle.get_scheduled_maintenance_type_display()} for {vehicle.scheduled_maintenance_date}: {vehicle.scheduled_maintenance_description}",
        )
        messages.success(
            request,
            f"{vehicle.get_scheduled_maintenance_type_display()} scheduled for {vehicle.scheduled_maintenance_date.strftime('%b %d, %Y')}.",
        )
        return

    messages.error(request, "Unsupported mechanic action. Use the fault report or maintenance checklist.")


def seacraft_operator_filter():
    return Q(license_authority='MARINA') | Q(license_number__istartswith='MAR-')


def vehicle_type_matches_operator(vehicle, driver):
    is_seacraft = vehicle.vehicle_type.name.strip().casefold() in {'marine', 'maritime'}
    is_seacraft_operator = (
        driver.license_authority == 'MARINA'
        or (driver.license_number or '').strip().upper().startswith('MAR-')
    )
    return is_seacraft == is_seacraft_operator


def assign_vehicle_operator(vehicle, driver):
    previous_vehicle = Vehicle.objects.filter(assigned_driver=driver).exclude(pk=vehicle.pk).first()
    if previous_vehicle:
        if previous_vehicle.status not in {'MAINTENANCE', 'PENDING_DISPOSAL'}:
            return previous_vehicle
        previous_vehicle.assigned_driver = None
        previous_vehicle.save(update_fields=['assigned_driver'])

    vehicle.assigned_driver = driver
    return None


# =========================================================================
# 1. HOMEPAGE VIEW
# =========================================================================
def homepage(request):
    all_vehicles = Vehicle.objects.all().select_related('vehicle_type', 'assigned_driver')
    sea_crafts = all_vehicles.filter(seacraft_vehicle_type_filter())
    land_vehicles = all_vehicles.exclude(seacraft_vehicle_type_filter())
    total_fleet = Vehicle.objects.exclude(status='disposal').count()
    land_operational_count = land_vehicles.filter(status='OPERATIONAL').count()
    sea_operational_count = sea_crafts.filter(status='OPERATIONAL').count()
    land_deployed_count = land_vehicles.filter(status='DEPLOYED').count()
    sea_deployed_count = sea_crafts.filter(status='DEPLOYED').count()
    land_maintenance_count = land_vehicles.filter(status='MAINTENANCE').count()
    sea_maintenance_count = sea_crafts.filter(status='MAINTENANCE').count()

    context = {
        'total_fleet': total_fleet,
        'sea_crafts': sea_crafts,
        'land_vehicles': land_vehicles,
        'land_operational_count': land_operational_count,
        'sea_operational_count': sea_operational_count,
        'land_deployed_count': land_deployed_count,
        'sea_deployed_count': sea_deployed_count,
        'land_maintenance_count': land_maintenance_count,
        'sea_maintenance_count': sea_maintenance_count,
        'total_count': all_vehicles.count(),
        'operational_count': all_vehicles.filter(status='OPERATIONAL').count(),
        'maintenance_count': all_vehicles.filter(status='MAINTENANCE').count(),
        'deployed_count': all_vehicles.filter(status='DEPLOYED').count(),
    }
    return render(request, 'fleet/homepage.html', context)


def deployment_activity_log(request):
    vehicle_content_type = ContentType.objects.get_for_model(Vehicle)
    deployment_events = LogEntry.objects.filter(
        content_type=vehicle_content_type,
    ).filter(
        Q(change_message__startswith='Deployed asset unit to ')
        | Q(change_message='Returned asset unit back to operational depot storage.')
    ).select_related('user').order_by('-action_time', '-pk')

    paginator = Paginator(deployment_events, 25)
    page = paginator.get_page(request.GET.get('page'))
    page_events = list(page.object_list)
    vehicle_by_id = {
        str(vehicle.pk): vehicle
        for vehicle in Vehicle.objects.filter(
            pk__in=[event.object_id for event in page_events]
        ).select_related('vehicle_type')
    }
    for event in page_events:
        vehicle = vehicle_by_id.get(str(event.object_id))
        event.asset_division = (
            'sea'
            if vehicle and vehicle.vehicle_type.name.strip().casefold() in {'marine', 'maritime'}
            else 'land'
        )

    context = {
        'deployment_page': page,
        'deployment_events': page_events,
        'deployment_event_count': paginator.count,
        'deployment_page_signature': ','.join(str(event.pk) for event in page_events),
    }
    if request.headers.get('x-requested-with') == 'XMLHttpRequest':
        return JsonResponse({
            'rows': render_to_string(
                'fleet/_deployment_activity_rows.html',
                context,
                request=request,
            ),
            'pagination': render_to_string(
                'fleet/_deployment_activity_pagination.html',
                context,
                request=request,
            ),
            'count': paginator.count,
            'signature': context['deployment_page_signature'],
        })

    return render(request, 'fleet/deployment_activity_log.html', context)


def fleet_report_permissions(user):
    user_groups = set(user.groups.values_list('name', flat=True))
    dashboard_groups = {'Logistics Officers', 'Technicians', 'Maritime_Tech', 'Seacraft Dispatch'}
    has_dashboard_group = bool(user_groups & dashboard_groups)
    is_logistics = user.is_superuser or 'Logistics Officers' in user_groups
    is_land_mechanic = 'Technicians' in user_groups
    is_seacraft_technician = 'Maritime_Tech' in user_groups
    is_seacraft_dispatch = (
        user.is_superuser
        or 'Seacraft Dispatch' in user_groups
    )
    if not has_dashboard_group and not user.is_superuser:
        is_logistics = check_user_role(user, 'Logistics Officers', ['logistics', 'log', 'depot', 'fleet'])
        is_land_mechanic = check_user_role(
            user,
            'Technicians',
            ['tech', 'technician', 'repair', 'maintenance', 'mechanic'],
        )
        is_seacraft_technician = check_user_role(
            user,
            'Maritime_Tech',
            ['maritime', 'sea', 'craft', 'tech', 'dispatch'],
        )
        is_seacraft_dispatch = any(
            keyword in user.username.lower() for keyword in ['sea', 'maritime']
        )
    if not (is_logistics or is_land_mechanic or is_seacraft_technician or is_seacraft_dispatch):
        return None

    allowed_divisions = set()
    if is_land_mechanic:
        allowed_divisions.add('Land Asset')
    if is_seacraft_technician or is_seacraft_dispatch:
        allowed_divisions.add('Seacraft')

    return {
        'is_logistics': is_logistics,
        'is_land_mechanic': is_land_mechanic,
        'is_seacraft_technician': is_seacraft_technician,
        'is_seacraft_dispatch': is_seacraft_dispatch,
        'allowed_divisions': allowed_divisions,
    }


def fleet_report_dashboard_name(permissions):
    if permissions['is_logistics']:
        return 'dashboard_portal:logistics_dashboard'
    if permissions['is_seacraft_dispatch']:
        return 'dashboard_portal:seacraft_dispatch'
    if permissions['is_seacraft_technician']:
        return 'dashboard_portal:seacraft_dashboard'
    return 'dashboard_portal:repairman_dashboard'


def last_assigned_driver_name(vehicle):
    if vehicle.assigned_driver_id and vehicle.assigned_driver:
        return vehicle.assigned_driver.name

    vehicle_content_type = ContentType.objects.get_for_model(Vehicle)
    deployment_message = LogEntry.objects.filter(
        content_type=vehicle_content_type,
        object_id=str(vehicle.pk),
        change_message__startswith='Deployed asset unit to ',
    ).order_by('-action_time', '-pk').values_list('change_message', flat=True).first()
    if not deployment_message or ' with operator ' not in deployment_message:
        return ''
    return deployment_message.rpartition(' with operator ')[2].rstrip('.')


@login_required
def fleet_report_incident(request):
    if request.method != 'POST':
        return HttpResponse("Incident reports must be submitted using the form.", status=405)

    permissions = fleet_report_permissions(request.user)
    if permissions is None:
        messages.error(request, "Access restricted to authorized fleet operations and maintenance accounts.")
        return redirect('dashboard_portal:homepage')

    vehicle_queryset = Vehicle.objects.select_related('vehicle_type')
    if not permissions['is_logistics']:
        division_filter = Q(pk__in=[])
        if 'Land Asset' in permissions['allowed_divisions']:
            division_filter |= ~seacraft_vehicle_type_filter()
        if 'Seacraft' in permissions['allowed_divisions']:
            division_filter |= seacraft_vehicle_type_filter()
        vehicle_queryset = vehicle_queryset.filter(division_filter)

    form = FleetIncidentForm(request.POST, vehicle_queryset=vehicle_queryset)
    if form.is_valid():
        incident = form.save(commit=False)
        incident.reported_by = request.user
        incident.last_assigned_driver = last_assigned_driver_name(incident.vehicle)
        incident.save()
        log_action_to_admin(
            request,
            incident,
            ADDITION,
            f"Reported {incident.get_incident_type_display().lower()} for {incident.vehicle}.",
        )
        messages.success(request, "Incident report saved and added to fleet incident reports.")
    else:
        for field, errors in form.errors.items():
            field_label = form.fields[field].label if field in form.fields else "Incident report"
            for error in errors:
                messages.error(request, f"{field_label}: {error}")
        if not form.errors:
            messages.error(request, "Check the incident details and try again.")

    return redirect(fleet_report_dashboard_name(permissions))


@login_required
def logistics_generate_report(request):
    permissions = fleet_report_permissions(request.user)
    if permissions is None:
        messages.error(request, "Access restricted to authorized fleet operations and maintenance accounts.")
        return redirect('dashboard_portal:homepage')

    is_logistics = permissions['is_logistics']
    allowed_divisions = permissions['allowed_divisions']

    report_format = request.GET.get('format', '').lower()
    report_kind = request.GET.get('report_kind', 'activity').lower()
    maintenance_category = request.GET.get('maintenance_category', 'all').lower()
    if report_kind not in {'activity', 'maintenance', 'damage', 'incident'}:
        return HttpResponse("Choose an activity, maintenance, damage, or incident report.", status=400)
    if maintenance_category not in {'all', 'regular', 'other'}:
        return HttpResponse("Choose all, regular, or other maintenance.", status=400)

    if report_kind == 'maintenance':
        report_title = 'Fleet Maintenance Report'
        report_heading = 'Maintenance Activity Report'
    elif report_kind == 'damage':
        report_title = 'Fleet Damage and Fault Report'
        report_heading = 'Damage and Fault Report'
    elif report_kind == 'incident':
        report_title = 'Fleet Incident Report'
        report_heading = 'Incident Report'
    else:
        report_title = 'PDRRMO Fleet Deployment Report'
        report_heading = 'Deployment Activity Report'

    period = request.GET.get('period', '')
    try:
        if period == 'month':
            selected = datetime.strptime(request.GET.get('month', ''), '%Y-%m').date()
            start_date = selected.replace(day=1)
            next_month = (start_date.replace(day=28) + timedelta(days=4)).replace(day=1)
            end_date = next_month - timedelta(days=1)
            period_label = start_date.strftime('%B %Y')
        elif period == 'week':
            year, week = request.GET.get('week', '').split('-W')
            start_date = date.fromisocalendar(int(year), int(week), 1)
            end_date = start_date + timedelta(days=6)
            period_label = f"Week {int(week):02d}, {year}"
        elif period == 'day':
            start_date = date.fromisoformat(request.GET.get('day', ''))
            end_date = start_date
            period_label = start_date.strftime('%B %d, %Y')
        elif period == 'custom':
            start_date = date.fromisoformat(request.GET.get('start_date', ''))
            end_date = date.fromisoformat(request.GET.get('end_date', ''))
            if start_date > end_date:
                raise ValueError("Start date must be on or before end date.")
            period_label = f"{start_date:%b %d, %Y} to {end_date:%b %d, %Y}"
        else:
            raise ValueError("Choose a report period.")
    except (TypeError, ValueError) as error:
        return HttpResponse(f"Invalid report period: {escape(str(error))}", status=400)

    if report_format not in {'print', 'csv', 'xlsx', 'pdf'}:
        return HttpResponse("Choose Print, CSV, Excel, or PDF as the export format.", status=400)

    start_datetime = timezone.make_aware(
        datetime.combine(start_date, time.min),
        timezone.get_current_timezone(),
    )
    end_datetime = timezone.make_aware(
        datetime.combine(end_date + timedelta(days=1), time.min),
        timezone.get_current_timezone(),
    )
    vehicle_content_type = ContentType.objects.get_for_model(Vehicle)
    event_query = LogEntry.objects.filter(
        content_type=vehicle_content_type,
        action_time__gte=start_datetime,
        action_time__lt=end_datetime,
    )
    rows = []
    incident_records = []
    if report_kind == 'incident':
        incident_query = FleetIncident.objects.filter(
            occurred_at__gte=start_datetime,
            occurred_at__lt=end_datetime,
        ).select_related('vehicle__vehicle_type', 'reported_by')
        incident_id = request.GET.get('incident_id')
        if incident_id:
            try:
                incident_id = int(incident_id)
                if incident_id < 1:
                    raise ValueError
            except ValueError:
                return HttpResponse("Choose a valid incident report.", status=400)
            incident_query = incident_query.filter(pk=incident_id)
        if not is_logistics:
            incident_division_filter = Q(pk__in=[])
            if 'Land Asset' in allowed_divisions:
                incident_division_filter |= ~seacraft_vehicle_type_filter()
            if 'Seacraft' in allowed_divisions:
                incident_division_filter |= seacraft_vehicle_type_filter()
            incident_query = incident_query.filter(vehicle__in=Vehicle.objects.filter(incident_division_filter))
        incident_records = list(incident_query.order_by('-occurred_at', '-pk'))
        for incident in incident_records:
            division = (
                'Seacraft'
                if incident.vehicle.vehicle_type.name.strip().casefold() in {'marine', 'maritime'}
                else 'Land Asset'
            )
            details = (
                f"{incident.get_incident_type_display()} — {incident.location}: "
                f"{incident.description}"
            )
            details += " | " + " | ".join((
                f"Injuries: {incident.get_injury_status_display()}",
                f"Injury / medical details: {incident.injury_details or 'Not recorded'}",
                f"Damage: {incident.damage_details or 'None reported'}",
                f"Witnesses: {incident.witnesses or 'Not recorded'}",
                f"Immediate actions: {incident.actions_taken or 'Not recorded'}",
                f"Follow-up recommendations: {incident.follow_up_recommendations or 'Not recorded'}",
            ))
            rows.append([
                timezone.localtime(incident.occurred_at).strftime('%Y-%m-%d %H:%M'),
                incident.get_incident_type_display(),
                division,
                str(incident.vehicle),
                details,
                incident.reported_by.username if incident.reported_by else 'System',
                incident.last_assigned_driver or 'Not recorded',
            ])
        event_query = event_query.none()
    elif report_kind == 'damage':
        event_query = event_query.filter(
            change_message__startswith='Reported maintenance fault for '
        )
    elif report_kind == 'maintenance':
        maintenance_events = (
            Q(change_message__startswith='Scheduled maintenance due: ')
            | Q(change_message__startswith='Scheduled Tire replacement for ')
            | Q(change_message__startswith='Scheduled Oil replacement for ')
            | Q(change_message__startswith='Scheduled Other maintenance for ')
            | Q(change_message__startswith='Reported maintenance fault for ')
            | Q(change_message__startswith='Completed maintenance checklist and returned ')
        )
        if maintenance_category == 'regular':
            maintenance_events &= (
                Q(change_message__startswith='Scheduled maintenance due: Tire replacement')
                | Q(change_message__startswith='Scheduled maintenance due: Oil replacement')
                | Q(change_message__startswith='Scheduled Tire replacement for ')
                | Q(change_message__startswith='Scheduled Oil replacement for ')
            )
        elif maintenance_category == 'other':
            maintenance_events &= (
                Q(change_message__startswith='Scheduled maintenance due: Other maintenance')
                | Q(change_message__startswith='Scheduled Other maintenance for ')
                | Q(change_message__startswith='Reported maintenance fault for ')
                | Q(change_message__startswith='Completed maintenance checklist and returned ')
            )
        event_query = event_query.filter(maintenance_events)
        category_label = {
            'all': 'All maintenance',
            'regular': 'Regular maintenance (tires and oil)',
            'other': 'Other maintenance and repairs',
        }[maintenance_category]
        report_title = f'Fleet Maintenance Report - {category_label}'
        report_heading = f'{category_label} Report'
    else:
        event_query = event_query.filter(
            Q(change_message__startswith='Deployed asset unit to ')
            | Q(change_message='Returned asset unit back to operational depot storage.')
        )
    events = list(event_query.select_related('user').order_by('-action_time', '-pk'))
    vehicles_by_id = {
        str(vehicle.pk): vehicle
        for vehicle in Vehicle.objects.filter(
            pk__in=[event.object_id for event in events]
        ).select_related('vehicle_type')
    }
    for event in events:
        vehicle = vehicles_by_id.get(str(event.object_id))
        division = 'Seacraft' if (
            vehicle and vehicle.vehicle_type.name.strip().casefold() in {'marine', 'maritime'}
        ) else 'Land Asset'
        if not is_logistics and division not in allowed_divisions:
            continue
        if report_kind == 'maintenance':
            if event.change_message.startswith('Scheduled maintenance due: '):
                activity = 'Scheduled maintenance due'
            elif event.change_message.startswith('Scheduled '):
                activity = 'Maintenance scheduled'
            elif event.change_message.startswith('Reported maintenance fault for '):
                activity = 'Fault reported'
            else:
                activity = 'Maintenance completed'
        elif report_kind == 'damage':
            activity = 'Damage / fault reported'
        else:
            activity = 'Deployment' if event.change_message.startswith('Deployed') else 'Return'
        rows.append([
            timezone.localtime(event.action_time).strftime('%Y-%m-%d %H:%M'),
            activity,
            division,
            event.object_repr,
            event.change_message,
            event.user.username if event.user else 'System',
            '',
        ])

    safe_period = f"{start_date:%Y%m%d}-{end_date:%Y%m%d}"
    headers = [
        'Date & Time',
        'Activity',
        'Division',
        'Fleet Asset',
        'Activity Details',
        'Recorded By',
        'Last Assigned Driver',
    ]
    spreadsheet_safe_rows = [
        [
            "'" + value if isinstance(value, str) and value.startswith(('=', '+', '-', '@')) else value
            for value in row
        ]
        for row in rows
    ]

    if report_format == 'print':
        if report_kind == 'incident':
            return render(
                request,
                'fleet/fleet_incident_report_print.html',
                {
                    'period_label': period_label,
                    'incident_records': incident_records,
                    'generated_at': timezone.localtime(),
                },
            )
        return render(
            request,
            'fleet/logistics_deployment_report_print.html',
            {
                'period_label': period_label,
                'headers': headers,
                'report_rows': rows,
                'generated_at': timezone.localtime(),
                'report_heading': report_heading,
                'report_title': report_title,
            },
        )

    if report_format == 'csv':
        output = io.StringIO(newline='')
        writer = csv.writer(output)
        writer.writerow([report_title, period_label])
        writer.writerow(headers)
        writer.writerows(spreadsheet_safe_rows)
        response = HttpResponse(output.getvalue(), content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = f'attachment; filename="fleet-{report_kind}-{safe_period}.csv"'
        return response

    if report_format == 'xlsx':
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Fleet Report'
        sheet.append([report_title, period_label])
        sheet.append(headers)
        for row in spreadsheet_safe_rows:
            sheet.append(row)
        sheet.freeze_panes = 'A3'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF', size=14)
            cell.fill = PatternFill('solid', fgColor='1A365D')
        for cell in sheet[2]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='2563EB')
        for column_index, width in enumerate((20, 20, 16, 34, 90, 24, 24), start=1):
            sheet.column_dimensions[get_column_letter(column_index)].width = width
        output = io.BytesIO()
        workbook.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = f'attachment; filename="fleet-{report_kind}-{safe_period}.xlsx"'
        return response

    output = io.BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(letter),
        rightMargin=0.4 * inch,
        leftMargin=0.4 * inch,
        topMargin=0.45 * inch,
        bottomMargin=0.45 * inch,
        title=f'{report_title} - {period_label}',
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'ReportTitle',
        parent=styles['Title'],
        textColor=colors.HexColor('#1A365D'),
        alignment=TA_CENTER,
        spaceAfter=6,
    )
    body_style = ParagraphStyle(
        'ReportCell',
        parent=styles['BodyText'],
        fontSize=7,
        leading=9,
    )
    report_data = [[Paragraph(f'<b>{escape(value)}</b>', body_style) for value in headers]]
    report_data.extend([
        [Paragraph(escape(str(value)), body_style) for value in row]
        for row in rows
    ])
    table = Table(
        report_data,
        repeatRows=1,
        colWidths=[0.9 * inch, 0.85 * inch, 0.7 * inch, 1.15 * inch, 4.3 * inch, 0.8 * inch, 1.0 * inch],
    )
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1A365D')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.35, colors.HexColor('#CBD5E1')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F1F5F9')]),
        ('LEFTPADDING', (0, 0), (-1, -1), 5),
        ('RIGHTPADDING', (0, 0), (-1, -1), 5),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    story = [
        Paragraph(escape(report_title), title_style),
        Paragraph(escape(period_label), styles['Heading3']),
        Spacer(1, 10),
        table,
    ]
    document.build(story)
    response = HttpResponse(output.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="fleet-{report_kind}-{safe_period}.pdf"'
    return response


# =========================================================================
# 2. CENTRAL ROUTER VIEW
# =========================================================================
@login_required
def dashboard_router(request):
    user = request.user
    user_group_names = list(user.groups.values_list('name', flat=True))
    user_groups_lower = [g.lower() for g in user_group_names]

    if user.is_superuser or 'superusers' in user_groups_lower:
        return redirect('/admin/')

    group_routing_matrix = {
        'Seacraft Dispatch':     'dashboard_portal:seacraft_dispatch',
        'Maritime_Tech':         'dashboard_portal:seacraft_dashboard',
        'Logistics Officers':    'dashboard_portal:logistics_dashboard',
        'Technicians':           'dashboard_portal:repairman_dashboard',
    }

    for group_name, destination_url in group_routing_matrix.items():
        if group_name in user_group_names:
            return redirect(destination_url)

    for group in user_groups_lower:
        if 'sea' in group or 'craft' in group or 'dispatch' in group:
            return redirect('dashboard_portal:seacraft_dispatch')
        elif 'tech' in group or 'nicians' in group or 'mechanic' in group:
            return redirect('dashboard_portal:repairman_dashboard')
        elif 'log' in group or 'depot' in group or 'manager' in group:
            return redirect('dashboard_portal:logistics_dashboard')

    if user_group_names:
        messages.info(request, f"Welcome {user.username}. Accessing general operations feed.")
        try:
            return redirect('dashboard_portal:homepage')
        except Exception:
            pass 

    return render(request, 'fleet/unassigned_pending.html')


# =========================================================================
# 3. REPAIRMAN DASHBOARD
# =========================================================================
@login_required
def repairman_dashboard(request):
    user = request.user
    is_authorized = check_user_role(
        user,
        'Technicians',
        ['tech', 'technician', 'repair', 'maintenance', 'mechanic']
    )
    
    if not is_authorized:
        messages.error(request, "Access restricted to authorized Repair Technicians.")
        return redirect('dashboard_portal:homepage')

    move_overdue_vehicles_to_maintenance(request)

    if request.method == 'POST':
        vehicle_id = request.POST.get('vehicle_id')
        action_type = request.POST.get('action_type')
        
        if vehicle_id:
            vehicle = get_object_or_404(
                Vehicle.objects.exclude(seacraft_vehicle_type_filter()),
                id=vehicle_id,
            )
            
            if action_type == 'FLAG_DISPOSAL':
                if vehicle.status != 'MAINTENANCE':
                    messages.error(request, f"{vehicle.model_name} must be in maintenance before it can be flagged for disposal.")
                    return redirect('dashboard_portal:repairman_dashboard')

                remarks = request.POST.get('disposal_remarks', '').strip()
                if not remarks:
                    messages.error(request, f"Failure: You must provide a maintenance justification remark to flag {vehicle.model_name} for disposal.")
                    return redirect('dashboard_portal:repairman_dashboard')

                vehicle.status = 'PENDING_DISPOSAL'
                if vehicle.assigned_driver:
                    vehicle.assigned_driver = None
                vehicle.save()
                
                log_action_to_admin(request, vehicle, CHANGE, f"Flagged asset {vehicle.model_name} for disposal from repair workshop. Remarks: {remarks}")
                messages.warning(request, f"{vehicle.model_name} has been routed to Logistics for decommissioning evaluation.")

            elif action_type == 'UNDO_DISPOSAL':
                if vehicle.status != 'PENDING_DISPOSAL':
                    messages.error(request, f"{vehicle.model_name} is not pending disposal.")
                    return redirect('dashboard_portal:repairman_dashboard')

                vehicle.status = 'MAINTENANCE'
                vehicle.save()
                log_action_to_admin(request, vehicle, CHANGE, f"Cancelled disposal request for {vehicle.model_name}; returned to maintenance by workshop.")
                messages.success(request, f"Disposal request for {vehicle.model_name} cancelled; returned to MAINTENANCE.")

            elif action_type == 'UPDATE_DISPOSAL_REASON':
                if vehicle.status != 'PENDING_DISPOSAL':
                    messages.error(request, f"{vehicle.model_name} is not pending disposal.")
                    return redirect('dashboard_portal:repairman_dashboard')

                remarks = request.POST.get('disposal_remarks', '').strip()
                if not remarks:
                    messages.error(request, "A disposal reason is required.")
                    return redirect('dashboard_portal:repairman_dashboard')

                log_action_to_admin(request, vehicle, CHANGE, f"Updated disposal reason for {vehicle.model_name}. Remarks: {remarks}")
                messages.success(request, f"Disposal reason updated for {vehicle.model_name}.")
            
            else:
                handle_mechanic_maintenance_action(request, vehicle)
                    
            return redirect('dashboard_portal:repairman_dashboard')

    vehicles = Vehicle.objects.exclude(
        seacraft_vehicle_type_filter()
    ).exclude(
        status__in=['ARCHIVED', 'DISPOSED']
    ).select_related(
        'vehicle_type', 'assigned_driver'
    ).order_by(
        Case(
            When(status='OPERATIONAL', then=Value(1)),
            When(status='DEPLOYED', then=Value(2)),
            When(status='MAINTENANCE', then=Value(3)),
            When(status='PENDING_DISPOSAL', then=Value(4)),
            default=Value(5),
            output_field=IntegerField(),
        ),
        'maintenance_started_at',
        'model_name'
    )
    vehicles = add_disposal_reasons(vehicles)
    
    for vehicle in vehicles:
        vehicle.fault_form = MaintenanceFaultForm(prefix=f'fault-{vehicle.pk}')
        vehicle.checklist_form = MaintenanceChecklistForm(
            checklist=LAND_MAINTENANCE_CHECKLIST,
            prefix=f'checklist-{vehicle.pk}',
        )
        vehicle.schedule_form = ScheduledMaintenanceForm(prefix=f'schedule-{vehicle.pk}')

    return render(request, 'fleet/repairman_dashboard.html', {
        'vehicles': vehicles,
        'incident_form': FleetIncidentForm(
            vehicle_queryset=Vehicle.objects.exclude(seacraft_vehicle_type_filter()).select_related('vehicle_type'),
        ),
        'maintenance_count': sum(
            vehicle.status in {'MAINTENANCE', 'PENDING_DISPOSAL'}
            for vehicle in vehicles
        ),
        'operational_count': sum(
            vehicle.status in {'OPERATIONAL', 'DEPLOYED'}
            for vehicle in vehicles
        ),
    })


# =========================================================================
# 4. SEACRAFT DASHBOARD
# =========================================================================
@login_required
def seacraft_dashboard(request):
    user = request.user
    is_authorized = check_user_role(
        user,
        'Maritime_Tech',
        ['maritime', 'sea', 'craft', 'tech', 'dispatch']
    )

    if not is_authorized:
        messages.error(request, "Access restricted to authorized Maritime Operators.")
        return redirect('dashboard_portal:homepage')

    move_overdue_vehicles_to_maintenance(request)

    if request.method == 'POST':
        vehicle_id = request.POST.get('vehicle_id')
        action_type = request.POST.get('action_type')
        
        if vehicle_id:
            vehicle = get_object_or_404(
                Vehicle.objects.filter(seacraft_vehicle_type_filter()),
                id=vehicle_id,
            )
            
            if action_type == 'FLAG_DISPOSAL':
                if vehicle.status != 'MAINTENANCE':
                    messages.error(request, f"{vehicle.model_name} must be in maintenance before it can be flagged for disposal.")
                    return redirect('dashboard_portal:seacraft_dashboard')

                remarks = request.POST.get('disposal_remarks', '').strip()
                if not remarks:
                    messages.error(request, f"Failure: You must provide a justification remark to flag {vehicle.model_name} for disposal.")
                    return redirect('dashboard_portal:seacraft_dashboard')

                vehicle.status = 'PENDING_DISPOSAL'
                if vehicle.assigned_driver:
                    vehicle.assigned_driver = None
                vehicle.save()
                
                log_action_to_admin(request, vehicle, CHANGE, f"Flagged maritime asset {vehicle.model_name} for disposal processing. Remarks: {remarks}")
                messages.warning(request, f"{vehicle.model_name} has been routed to Logistics for disposal confirmation.")

            elif action_type == 'UNDO_DISPOSAL':
                if vehicle.status != 'PENDING_DISPOSAL':
                    messages.error(request, f"{vehicle.model_name} is not pending disposal.")
                    return redirect('dashboard_portal:seacraft_dashboard')

                vehicle.status = 'MAINTENANCE'
                vehicle.save()
                log_action_to_admin(request, vehicle, CHANGE, f"Cancelled disposal request for {vehicle.model_name}; returned to maintenance by maritime workshop.")
                messages.success(request, f"Disposal request for {vehicle.model_name} cancelled; returned to MAINTENANCE.")

            elif action_type == 'UPDATE_DISPOSAL_REASON':
                if vehicle.status != 'PENDING_DISPOSAL':
                    messages.error(request, f"{vehicle.model_name} is not pending disposal.")
                    return redirect('dashboard_portal:seacraft_dashboard')

                remarks = request.POST.get('disposal_remarks', '').strip()
                if not remarks:
                    messages.error(request, "A disposal reason is required.")
                    return redirect('dashboard_portal:seacraft_dashboard')

                log_action_to_admin(request, vehicle, CHANGE, f"Updated disposal reason for {vehicle.model_name}. Remarks: {remarks}")
                messages.success(request, f"Disposal reason updated for {vehicle.model_name}.")
            
            else:
                handle_mechanic_maintenance_action(request, vehicle)
            
            return redirect('dashboard_portal:seacraft_dashboard')

    vehicles = Vehicle.objects.filter(
        seacraft_vehicle_type_filter()
    ).exclude(
        status__in=['ARCHIVED', 'DISPOSED']
    ).select_related(
        'vehicle_type', 'assigned_driver'
    ).order_by(
        Case(
            When(status='OPERATIONAL', then=Value(1)),
            When(status='DEPLOYED', then=Value(2)),
            When(status='MAINTENANCE', then=Value(3)),
            When(status='PENDING_DISPOSAL', then=Value(4)),
            default=Value(5),
            output_field=IntegerField(),
        ),
        'maintenance_started_at',
        'model_name'
    )
    vehicles = add_disposal_reasons(vehicles)

    for vehicle in vehicles:
        vehicle.fault_form = MaintenanceFaultForm(prefix=f'fault-{vehicle.pk}')
        vehicle.checklist_form = MaintenanceChecklistForm(
            checklist=SEACRAFT_MAINTENANCE_CHECKLIST,
            prefix=f'checklist-{vehicle.pk}',
        )
        vehicle.schedule_form = ScheduledMaintenanceForm(prefix=f'schedule-{vehicle.pk}')

    return render(request, 'fleet/seacraft_dashboard.html', {
        'vehicles': vehicles,
        'incident_form': FleetIncidentForm(
            vehicle_queryset=Vehicle.objects.filter(seacraft_vehicle_type_filter()).select_related('vehicle_type'),
        ),
        'maintenance_count': sum(
            vehicle.status in {'MAINTENANCE', 'PENDING_DISPOSAL'}
            for vehicle in vehicles
        ),
        'operational_count': sum(
            vehicle.status in {'OPERATIONAL', 'DEPLOYED'}
            for vehicle in vehicles
        ),
    })


# =========================================================================
# 5. LOGISTICS DASHBOARD
# =========================================================================
@login_required
def logistics_dashboard(request):
    user = request.user

    if not check_user_role(user, 'Logistics Officers', ['logistics', 'log', 'depot', 'fleet']):
        messages.error(request, "Access restricted to Logistics Depot management accounts.")
        return redirect('dashboard_portal:homepage')

    move_overdue_vehicles_to_maintenance(request)

    if request.method == 'POST':
        action = request.POST.get('action')
        vehicle_id = request.POST.get('vehicle_id')

        if action == 'add_operator':
            form = OperatorDetailsForm(request.POST)
            if form.is_valid():
                operator = form.save()
                log_action_to_admin(
                    request,
                    operator,
                    ADDITION,
                    f"Registered {form.cleaned_data['operator_type'].lower()} operator {operator.name}.",
                )
                messages.success(request, f"Operator details for {operator.name} were added.")
            else:
                for errors in form.errors.values():
                    for error in errors:
                        messages.error(request, error)
            return redirect('dashboard_portal:logistics_dashboard')

        # ACTION A: DEPLOY VEHICLE (Requires Location, Purpose, Time, & Operator Validation)
        if action == 'deploy_vehicle':
            vehicle = get_object_or_404(Vehicle, id=vehicle_id)
            if vehicle.status != 'OPERATIONAL':
                messages.error(request, f"'{vehicle.model_name}' must be operational before it can be deployed.")
                return redirect('dashboard_portal:logistics_dashboard')
            
            location = request.POST.get('deployment_location', '').strip()
            purpose = request.POST.get('deployment_purpose', '').strip()
            deploy_time_str = request.POST.get('deployment_time', '').strip()
            driver_id = request.POST.get('driver_id', '').strip()

            if not location or not purpose or not deploy_time_str:
                messages.error(request, "Deployment location, deployment purpose, and deployment time are required to dispatch an asset.")
                return redirect('dashboard_portal:logistics_dashboard')

            parsed_time = parse_datetime(deploy_time_str)
            if parsed_time is None:
                messages.error(request, "Enter a valid deployment date and time.")
                return redirect('dashboard_portal:logistics_dashboard')
            if timezone.is_naive(parsed_time):
                parsed_time = timezone.make_aware(parsed_time, timezone.get_current_timezone())

            if not driver_id:
                messages.error(request, f"DISPATCH DENIED: Personnel operator must be selected prior to deploying '{vehicle.model_name}'.")
                return redirect('dashboard_portal:logistics_dashboard')

            driver_obj = Driver.objects.filter(pk=driver_id, is_active=True).first()
            if driver_obj is None or not vehicle_type_matches_operator(vehicle, driver_obj):
                messages.error(request, "Select an active operator qualified for this vehicle type.")
                return redirect('dashboard_portal:logistics_dashboard')

            with transaction.atomic():
                conflicting_vehicle = assign_vehicle_operator(vehicle, driver_obj)
                if conflicting_vehicle:
                    messages.error(
                        request,
                        f"RESTRICTION: Personnel '{driver_obj.name}' is already assigned to active fleet asset '{conflicting_vehicle.model_name}'.",
                    )
                    return redirect('dashboard_portal:logistics_dashboard')
                vehicle.status = 'DEPLOYED'
                vehicle.deployment_location = location
                vehicle.deployment_purpose = purpose
                vehicle.deployment_time = parsed_time
                vehicle.save()

            log_action_to_admin(
                request, 
                vehicle, 
                CHANGE, 
                f"Deployed asset unit to {location} for '{purpose}' at {parsed_time.strftime('%Y-%m-%d %H:%M')} with operator {vehicle.assigned_driver.name}."
            )
            messages.success(request, f"Asset unit '{vehicle.model_name}' deployed to '{location}' for '{purpose}' with operator '{vehicle.assigned_driver.name}'!")
            return redirect('dashboard_portal:logistics_dashboard')

        # ACTION B: RETURN VEHICLE TO DEPOT BASE
        elif action == 'return_vehicle':
            vehicle = get_object_or_404(Vehicle, id=vehicle_id)
            if vehicle.status != 'DEPLOYED':
                messages.error(request, f"'{vehicle.model_name}' is not currently deployed.")
                return redirect('dashboard_portal:logistics_dashboard')
            vehicle.status = 'OPERATIONAL'
            vehicle.deployment_location = None
            vehicle.deployment_purpose = None
            vehicle.deployment_time = None
            vehicle.save()
            log_action_to_admin(request, vehicle, CHANGE, "Returned asset unit back to operational depot storage.")
            messages.success(request, f"Asset unit {vehicle.model_name} returned to depot!")
            return redirect('dashboard_portal:logistics_dashboard')

        # ACTION C: NEW ASSET REGISTRATION ONBOARDING
        elif action == 'add_vehicle':
            if 'asset_count' in request.POST:
                try:
                    asset_count = int(request.POST['asset_count'])
                except (TypeError, ValueError):
                    asset_count = 0
                if not 1 <= asset_count <= 20:
                    messages.error(request, "Add between 1 and 20 fleet assets at a time.")
                    return redirect('dashboard_portal:logistics_dashboard')
                registration_forms = [
                    FleetAssetRegistrationForm(request.POST, prefix=f'asset-{index}')
                    for index in range(asset_count)
                ]
            else:
                model_name = request.POST.get('model_name', '').strip()
                plate_number = request.POST.get('plate_number', '').strip()
                type_id = request.POST.get('vehicle_type', '').strip()
                if not model_name or not plate_number or not type_id:
                    messages.error(request, "Please enter the asset name, plate number or hull ID, and asset type.")
                    return redirect('dashboard_portal:logistics_dashboard')
                if Vehicle.objects.filter(plate_number=plate_number).exists():
                    messages.error(request, f"An asset with plate number or hull ID '{plate_number}' is already registered.")
                    return redirect('dashboard_portal:logistics_dashboard')
                registration_forms = [FleetAssetRegistrationForm(request.POST)]

            valid_forms = [form.is_valid() for form in registration_forms]
            forms_valid = all(valid_forms)
            submitted_identifiers = [
                form.cleaned_data.get('plate_number', '').strip()
                for form, is_valid in zip(registration_forms, valid_forms)
                if is_valid
            ]
            duplicate_identifiers = {
                identifier
                for identifier in submitted_identifiers
                if submitted_identifiers.count(identifier) > 1
            }
            if duplicate_identifiers:
                messages.error(
                    request,
                    "Each asset must have a different plate number or hull ID.",
                )
                forms_valid = False

            if forms_valid:
                for form in registration_forms:
                    identifier = form.cleaned_data['plate_number'].strip()
                    model_name = form.cleaned_data['model_name'].strip()
                    if Vehicle.objects.filter(plate_number=identifier).exists():
                        messages.error(request, f"An asset with plate number or hull ID '{identifier}' is already registered.")
                        forms_valid = False
                        break
                    if len(model_name) > 100 or len(identifier) > 50:
                        messages.error(request, "The asset name must be 100 characters or fewer and the plate number or hull ID must be 50 characters or fewer.")
                        forms_valid = False
                        break

            if not forms_valid:
                for index, form in enumerate(registration_forms, start=1):
                    for field, errors in form.errors.items():
                        field_label = form.fields[field].label
                        for error in errors:
                            messages.error(request, f"Asset {index} — {field_label}: {error}")
                return redirect('dashboard_portal:logistics_dashboard')

            registered_assets = []
            with transaction.atomic():
                for form in registration_forms:
                    new_asset = form.save(commit=False)
                    v_type = form.cleaned_data['vehicle_type']
                    new_asset.vehicle_type = v_type
                    new_asset.status = 'OPERATIONAL'
                    is_seacraft = v_type.name.strip().casefold() in {'marine', 'maritime'}
                    if is_seacraft:
                        new_asset.make = ''
                        new_asset.model_year = None
                        new_asset.color = ''
                        new_asset.odometer_km = None
                    else:
                        new_asset.hull_type = ''
                        new_asset.length_m = None
                    new_asset.save()
                    log_action_to_admin(request, new_asset, ADDITION, f"Registered new asset unit '{new_asset.model_name}' into inventory records.")
                    registered_assets.append(new_asset)
            if len(registered_assets) == 1:
                messages.success(request, f"New fleet asset '{registered_assets[0].model_name}' has been securely registered to the base depot map.")
            else:
                messages.success(request, f"{len(registered_assets)} fleet assets have been registered.")
            return redirect('dashboard_portal:logistics_dashboard')

        # ACTION D: SET/UNSET PERSONNEL OPERATOR
        elif action == 'set_driver':
            vehicle = get_object_or_404(Vehicle, id=vehicle_id)
            if vehicle.status in {'DEPLOYED', 'PENDING_DISPOSAL', 'ARCHIVED'}:
                messages.error(request, f"Operator assignments cannot be changed while '{vehicle.model_name}' is {vehicle.get_status_display().lower()}.")
                return redirect('dashboard_portal:logistics_dashboard')
            driver_id = request.POST.get('driver_id')
            
            if driver_id:  
                driver_obj = Driver.objects.filter(pk=driver_id, is_active=True).first()
                if driver_obj is None or not vehicle_type_matches_operator(vehicle, driver_obj):
                    messages.error(request, "Select an active operator qualified for this vehicle type.")
                    return redirect('dashboard_portal:logistics_dashboard')
                msg = f"Assigned operator {driver_obj.name} to {vehicle.model_name}."
            else:  
                vehicle.assigned_driver = None
                msg = f"Removed driver assignment from asset {vehicle.model_name}."
                
            with transaction.atomic():
                if driver_id:
                    conflicting_vehicle = assign_vehicle_operator(vehicle, driver_obj)
                    if conflicting_vehicle:
                        messages.error(
                            request,
                            f"RESTRICTION: Personnel '{driver_obj.name}' is already assigned to active fleet asset '{conflicting_vehicle.model_name}'.",
                        )
                        return redirect('dashboard_portal:logistics_dashboard')
                vehicle.save()
            log_action_to_admin(request, vehicle, CHANGE, msg)
            messages.success(request, msg)
            return redirect('dashboard_portal:logistics_dashboard')

        # ACTION E: CONFIRM DISPOSAL PIPELINE
        elif action == 'confirm_disposal':
            vessel_to_archive = get_object_or_404(Vehicle, id=vehicle_id)
            if vessel_to_archive.status != 'PENDING_DISPOSAL':
                messages.error(request, f"'{vessel_to_archive.model_name}' is not pending disposal.")
                return redirect('dashboard_portal:logistics_dashboard')
            vessel_to_archive.status = 'ARCHIVED'
            vessel_to_archive.assigned_driver = None
            vessel_to_archive.deployment_location = None
            vessel_to_archive.deployment_purpose = None
            vessel_to_archive.deployment_time = None
            vessel_to_archive.save()
            log_action_to_admin(request, vessel_to_archive, CHANGE, f"Approved and permanently archived asset: {vessel_to_archive.model_name}")
            messages.success(request, f"Asset {vessel_to_archive.model_name} successfully moved to secure archives.")
            return redirect('dashboard_portal:logistics_dashboard')

        # ACTION F: DECLINE DISPOSAL PIPELINE
        elif action == 'decline_disposal':
            vessel_to_repair = get_object_or_404(Vehicle, id=vehicle_id)
            if vessel_to_repair.status != 'PENDING_DISPOSAL':
                messages.error(request, f"'{vessel_to_repair.model_name}' is not pending disposal.")
                return redirect('dashboard_portal:logistics_dashboard')
            vessel_to_repair.status = 'MAINTENANCE'
            vessel_to_repair.save()
            log_action_to_admin(request, vessel_to_repair, CHANGE, f"Rejected disposal request. Returned to maintenance array: {vessel_to_repair.model_name}")
            messages.info(request, f"Disposal declined. {vessel_to_repair.model_name} reverted to MAINTENANCE status.")
            return redirect('dashboard_portal:logistics_dashboard')

        else:
            messages.error(request, "Unsupported logistics dashboard action.")
            return redirect('dashboard_portal:logistics_dashboard')
    all_vehicles = Vehicle.objects.all().select_related('vehicle_type', 'assigned_driver')
    
    status_order = Case(
        When(status='OPERATIONAL', then=Value(1)),
        When(status='DEPLOYED', then=Value(2)),
        When(status='MAINTENANCE', then=Value(3)),
        When(status='PENDING_DISPOSAL', then=Value(4)),
        default=Value(5),
        output_field=IntegerField(),
    )
    land_vehicles = all_vehicles.exclude(
        seacraft_vehicle_type_filter()
    ).exclude(
        status__in=['ARCHIVED', 'DISPOSED']
    ).order_by(status_order, 'model_name')
    sea_crafts = all_vehicles.filter(
        seacraft_vehicle_type_filter()
    ).exclude(
        status__in=['ARCHIVED', 'DISPOSED']
    ).order_by(status_order, 'model_name')

    land_vehicles = add_disposal_reasons(land_vehicles)
    sea_crafts = add_disposal_reasons(sea_crafts)

    land_standby = [v for v in land_vehicles if v.status == 'OPERATIONAL']
    land_deployed = [v for v in land_vehicles if v.status == 'DEPLOYED']
    land_maintenance = [v for v in land_vehicles if v.status in ['MAINTENANCE', 'PENDING_DISPOSAL']]

    sea_standby = [c for c in sea_crafts if c.status == 'OPERATIONAL']
    sea_deployed = [c for c in sea_crafts if c.status == 'DEPLOYED']
    sea_maintenance = [c for c in sea_crafts if c.status in ['MAINTENANCE', 'PENDING_DISPOSAL']]

    all_drivers = Driver.objects.filter(is_active=True)
    land_drivers = all_drivers.exclude(seacraft_operator_filter())
    sea_drivers = all_drivers.filter(seacraft_operator_filter())

    assigned_driver_ids = set(
        Vehicle.objects.filter(assigned_driver__isnull=False).values_list('assigned_driver_id', flat=True)
    )
    
    assigned_driver_map = {
        v.assigned_driver_id: v.model_name for v in Vehicle.objects.filter(assigned_driver__isnull=False)
    }

    types = VehicleType.objects.all()
    
    context = {
        'land_vehicles': land_vehicles, 
        'sea_crafts': sea_crafts, 
        'land_standby': land_standby,
        'land_deployed': land_deployed,
        'land_maintenance': land_maintenance,
        'sea_standby': sea_standby,
        'sea_deployed': sea_deployed,
        'sea_maintenance': sea_maintenance,
        'types': types, 
        'land_drivers': land_drivers,
        'sea_drivers': sea_drivers,
        'assigned_driver_ids': assigned_driver_ids,
        'assigned_driver_map': assigned_driver_map,
        'operator_form': OperatorDetailsForm(),
        'asset_registration_forms': [
            FleetAssetRegistrationForm(prefix='asset-0'),
        ],
        'incident_form': FleetIncidentForm(
            vehicle_queryset=all_vehicles.select_related('vehicle_type'),
        ),
        'incident_records': FleetIncident.objects.select_related(
            'vehicle__vehicle_type',
            'reported_by',
        ).order_by('-occurred_at', '-pk'),
    }
    return render(request, 'fleet/logistics_dashboard.html', context)


# =========================================================================
# 6. SEACRAFT DISPATCH VIEW
# =========================================================================
@login_required
def seacraft_dispatch_view(request):
    user = request.user
    if not (
        user.is_superuser
        or user.groups.filter(name="Seacraft Dispatch").exists()
        or any(x in user.username.lower() for x in ["sea", "maritime"])
    ):
        messages.error(request, "Access restricted to authorized Maritime Dispatchers.")
        return redirect("dashboard_portal:homepage")

    move_overdue_vehicles_to_maintenance(request)

    if request.method == "POST":
        action = request.POST.get("action") or request.POST.get("action_type")
        if action == "add_operator":
            form = OperatorDetailsForm(
                request.POST,
                allowed_operator_type='SEA',
            )
            if form.is_valid():
                operator = form.save()
                log_action_to_admin(
                    request,
                    operator,
                    ADDITION,
                    f"Registered seacraft operator {operator.name}.",
                )
                messages.success(request, f"Seacraft operator details for {operator.name} were added.")
            else:
                for errors in form.errors.values():
                    for error in errors:
                        messages.error(request, error)
            return redirect("dashboard_portal:seacraft_dispatch")

        if action == "add_seacraft":
            if 'craft_count' in request.POST:
                try:
                    craft_count = int(request.POST['craft_count'])
                except (TypeError, ValueError):
                    craft_count = 0
                if not 1 <= craft_count <= 20:
                    messages.error(request, "Add between 1 and 20 seacrafts at a time.")
                    return redirect("dashboard_portal:seacraft_dispatch")
                forms = [
                    SeacraftRegistrationForm(request.POST, prefix=f'craft-{index}')
                    for index in range(craft_count)
                ]
            else:
                forms = [SeacraftRegistrationForm(request.POST)]

            valid_forms = [form.is_valid() for form in forms]
            submitted_identifiers = [
                form.cleaned_data.get('plate_number', '').strip()
                for form, is_valid in zip(forms, valid_forms)
                if is_valid
            ]
            duplicate_identifiers = {
                identifier
                for identifier in submitted_identifiers
                if submitted_identifiers.count(identifier) > 1
            }
            if duplicate_identifiers:
                messages.error(request, "Each seacraft must have a different hull ID.")

            form_errors = any(not is_valid for is_valid in valid_forms)
            already_registered = any(
                Vehicle.objects.filter(plate_number=identifier).exists()
                for identifier in submitted_identifiers
            )
            if already_registered:
                messages.error(request, "A seacraft with one of those hull IDs is already registered.")
            if form_errors or duplicate_identifiers or already_registered:
                for index, form in enumerate(forms, start=1):
                    for field, errors in form.errors.items():
                        field_label = form.fields[field].label if field in form.fields else 'Seacraft'
                        for error in errors:
                            messages.error(request, f"Seacraft {index} — {field_label}: {error}")
                return redirect("dashboard_portal:seacraft_dispatch")

            registered_crafts = []
            with transaction.atomic():
                for form in forms:
                    new_craft = form.save(commit=False)
                    new_craft.status = "OPERATIONAL"
                    new_craft.save()
                    log_action_to_admin(
                        request,
                        new_craft,
                        ADDITION,
                        f"Registered new seacraft unit '{new_craft.model_name}' into inventory records.",
                    )
                    registered_crafts.append(new_craft)
            if len(registered_crafts) == 1:
                messages.success(request, f"Seacraft '{registered_crafts[0].model_name}' registered.")
            else:
                messages.success(request, f"{len(registered_crafts)} seacrafts registered.")
            return redirect("dashboard_portal:seacraft_dispatch")

        vehicle = get_object_or_404(
            Vehicle.objects.filter(seacraft_vehicle_type_filter()),
            id=request.POST.get("vehicle_id"),
        )

        if action == "FLAG_DISPOSAL":
            if vehicle.status != "MAINTENANCE":
                messages.error(request, f"{vehicle.model_name} must be in maintenance before it can be flagged for disposal.")
                return redirect("dashboard_portal:seacraft_dispatch")
            remarks = request.POST.get("disposal_remarks", "").strip()
            if not remarks:
                messages.error(request, f"Failure: You must provide a justification remark to flag {vehicle.model_name} for disposal.")
                return redirect("dashboard_portal:seacraft_dispatch")

            vehicle.status = "PENDING_DISPOSAL"
            vehicle.assigned_driver = None
            vehicle.deployment_location = None
            vehicle.deployment_purpose = None
            vehicle.deployment_time = None
            vehicle.save(update_fields=[
                "status",
                "assigned_driver",
                "deployment_location",
                "deployment_purpose",
                "deployment_time",
            ])

            log_action_to_admin(request, vehicle, CHANGE, f"Flagged maritime asset {vehicle.model_name} for disposal. Remarks: {remarks}")
            messages.warning(request, f"{vehicle.model_name} has been routed to Logistics for disposal confirmation.")

        elif action in ["DISPATCH_MISSION", "deploy_vehicle"]:
            if vehicle.status != "OPERATIONAL":
                messages.error(request, f"'{vehicle.model_name}' must be operational before it can be deployed.")
                return redirect("dashboard_portal:seacraft_dispatch")

            location = request.POST.get("deployment_location", "").strip()
            purpose = request.POST.get("deployment_purpose", "").strip()
            deploy_time_str = request.POST.get("deployment_time", "").strip()
            driver_id = request.POST.get("driver_id", "").strip()

            if not location or not purpose or not deploy_time_str:
                messages.error(request, "Deployment location, deployment purpose, and deployment time are required to dispatch a vessel.")
                return redirect("dashboard_portal:seacraft_dispatch")

            parsed_time = parse_datetime(deploy_time_str)
            if parsed_time is None:
                messages.error(request, "Enter a valid deployment date and time.")
                return redirect("dashboard_portal:seacraft_dispatch")
            if timezone.is_naive(parsed_time):
                parsed_time = timezone.make_aware(parsed_time, timezone.get_current_timezone())

            driver = Driver.objects.filter(
                pk=driver_id,
                is_active=True,
            ).filter(seacraft_operator_filter()).first()
            if driver is None:
                messages.error(request, "Select an active, qualified seacraft operator before deployment.")
                return redirect("dashboard_portal:seacraft_dispatch")

            with transaction.atomic():
                conflicting_vehicle = assign_vehicle_operator(vehicle, driver)
                if conflicting_vehicle:
                    messages.error(
                        request,
                        f"Personnel operator '{driver.name}' is assigned to active fleet asset '{conflicting_vehicle.model_name}'.",
                    )
                    return redirect("dashboard_portal:seacraft_dispatch")
                vehicle.status = "DEPLOYED"
                vehicle.deployment_location = location
                vehicle.deployment_purpose = purpose
                vehicle.deployment_time = parsed_time
                vehicle.save()

            log_action_to_admin(
                request,
                vehicle,
                CHANGE,
                f"Deployed asset unit to {location} for '{purpose}' at {parsed_time.strftime('%Y-%m-%d %H:%M')} with operator {driver.name}.",
            )
            messages.success(request, f"Vessel '{vehicle.model_name}' deployed to '{location}' for '{purpose}' with operator '{driver.name}'.")

        elif action == "set_driver":
            if vehicle.status in {"DEPLOYED", "PENDING_DISPOSAL", "ARCHIVED"}:
                messages.error(request, f"Operator assignments cannot be changed while '{vehicle.model_name}' is {vehicle.get_status_display().lower()}.")
                return redirect("dashboard_portal:seacraft_dispatch")
            driver_id = request.POST.get("driver_id", "").strip()
            if driver_id:
                driver = Driver.objects.filter(
                    pk=driver_id,
                    is_active=True,
                ).filter(seacraft_operator_filter()).first()
                if driver is None:
                    messages.error(request, "Select an active, qualified seacraft operator.")
                    return redirect("dashboard_portal:seacraft_dispatch")
                msg = f"Operator assignment updated for {vehicle.model_name} to {driver.name}."
            else:
                vehicle.assigned_driver = None
                msg = f"Removed operator assignment from asset {vehicle.model_name}."

            with transaction.atomic():
                if driver_id:
                    conflicting_vehicle = assign_vehicle_operator(vehicle, driver)
                    if conflicting_vehicle:
                        messages.error(
                            request,
                            f"Personnel operator '{driver.name}' is assigned to active fleet asset '{conflicting_vehicle.model_name}'.",
                        )
                        return redirect("dashboard_portal:seacraft_dispatch")
                vehicle.save()
            log_action_to_admin(request, vehicle, CHANGE, msg)
            messages.success(request, msg)

        elif action == "return_vehicle":
            if vehicle.status != "DEPLOYED":
                messages.error(request, f"'{vehicle.model_name}' is not currently deployed.")
                return redirect("dashboard_portal:seacraft_dispatch")
            vehicle.status = "OPERATIONAL"
            vehicle.deployment_location = None
            vehicle.deployment_purpose = None
            vehicle.deployment_time = None
            vehicle.save(update_fields=[
                "status",
                "deployment_location",
                "deployment_purpose",
                "deployment_time",
            ])
            log_action_to_admin(request, vehicle, CHANGE, "Returned asset unit back to operational depot storage.")
            messages.success(request, f"Vessel '{vehicle.model_name}' returned to depot.")

        elif action == "confirm_disposal":
            if vehicle.status != "PENDING_DISPOSAL":
                messages.error(request, f"'{vehicle.model_name}' is not pending disposal.")
                return redirect("dashboard_portal:seacraft_dispatch")
            vehicle.status = "ARCHIVED"
            vehicle.assigned_driver = None
            vehicle.deployment_location = None
            vehicle.deployment_purpose = None
            vehicle.deployment_time = None
            vehicle.save(update_fields=[
                "status",
                "assigned_driver",
                "deployment_location",
                "deployment_purpose",
                "deployment_time",
            ])
            log_action_to_admin(request, vehicle, CHANGE, f"Approved and permanently archived asset: {vehicle.model_name}")
            messages.success(request, f"Asset {vehicle.model_name} successfully moved to secure archives.")

        elif action == "decline_disposal":
            if vehicle.status != "PENDING_DISPOSAL":
                messages.error(request, f"'{vehicle.model_name}' is not pending disposal.")
                return redirect("dashboard_portal:seacraft_dispatch")
            vehicle.status = "MAINTENANCE"
            vehicle.save(update_fields=["status"])
            log_action_to_admin(request, vehicle, CHANGE, f"Rejected disposal request. Returned to maintenance array: {vehicle.model_name}")
            messages.info(request, f"Disposal declined. {vehicle.model_name} reverted to MAINTENANCE status.")

        else:
            messages.error(request, "Unsupported seacraft dispatch action.")

        return redirect("dashboard_portal:seacraft_dispatch")

    all_drivers = Driver.objects.filter(is_active=True)
    sea_drivers = all_drivers.filter(seacraft_operator_filter())

    assigned_driver_ids = set(
        Vehicle.objects.filter(assigned_driver__isnull=False).values_list('assigned_driver_id', flat=True)
    )

    sea_crafts = (
        Vehicle.objects.filter(seacraft_vehicle_type_filter())
        .exclude(status__in=["ARCHIVED", "DISPOSED"])
        .select_related("vehicle_type", "assigned_driver")
        .order_by(
            Case(
                When(status="OPERATIONAL", then=Value(1)),
                When(status="DEPLOYED", then=Value(2)),
                When(status="MAINTENANCE", then=Value(3)),
                When(status="PENDING_DISPOSAL", then=Value(4)),
                default=Value(5),
                output_field=IntegerField(),
            ),
            "model_name",
        )
    )
    sea_crafts = add_disposal_reasons(sea_crafts)
    sea_standby = [craft for craft in sea_crafts if craft.status == "OPERATIONAL"]
    sea_deployed = [craft for craft in sea_crafts if craft.status == "DEPLOYED"]
    sea_maintenance = [
        craft for craft in sea_crafts
        if craft.status in {"MAINTENANCE", "PENDING_DISPOSAL"}
    ]

    return render(
        request,
        "fleet/seacraft_dispatch.html",
        {
            "sea_crafts": sea_crafts,
            "sea_standby": sea_standby,
            "sea_deployed": sea_deployed,
            "sea_maintenance": sea_maintenance,
            "drivers": sea_drivers,
            "assigned_driver_ids": assigned_driver_ids,
            "operator_form": OperatorDetailsForm(allowed_operator_type='SEA'),
            "seacraft_forms": [SeacraftRegistrationForm(prefix='craft-0')],
            "incident_form": FleetIncidentForm(
                vehicle_queryset=sea_crafts.select_related('vehicle_type'),
            ),
        },
    )