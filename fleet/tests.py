import json
from datetime import datetime, timedelta
from io import BytesIO

from django.contrib.admin.models import CHANGE, LogEntry
from django.contrib.auth.models import Group, User
from django.contrib.contenttypes.models import ContentType
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from .models import Driver, FleetIncident, Vehicle, VehicleType


class LoginViewTests(TestCase):
    def test_login_page_renders_without_authenticating_on_get(self):
        response = self.client.get(reverse('dashboard_portal:login'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Fleet Operations Center')
        self.assertNotContains(response, 'ACCESS DENIED: Invalid Username or Password.')

    def test_login_page_shows_error_for_invalid_credentials(self):
        response = self.client.post(
            reverse('dashboard_portal:login'),
            {'username': 'unknown-user', 'password': 'wrong-password'},
            follow=True,
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'ACCESS DENIED: Invalid Username or Password.')

    def test_logistics_dashboard_allows_non_staff_user_in_logistics_group(self):
        group = Group.objects.create(name='Logistics Officers')
        user = User.objects.create_user(username='logistics_user', password='SecurePass123!')
        user.groups.add(group)
        user.is_staff = False
        user.save()

        self.client.force_login(user)
        response = self.client.get(reverse('dashboard_portal:logistics_dashboard'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'id="logisticsMenuDrawer"')
        self.assertContains(response, 'data-menu-action="operators"')
        self.assertContains(response, 'data-menu-action="add-operator"')
        self.assertContains(response, 'data-menu-action="add-asset"')
        self.assertContains(response, 'data-menu-action="print"')
        self.assertNotContains(response, 'id_license_authority')
        self.assertContains(response, 'Issuing authority is selected automatically from operator type.')
        self.assertContains(response, 'license authority: ${authority} (automatic).')
        self.assertContains(response, 'data-menu-action="generate-report"')
        self.assertContains(response, 'Generate Fleet Report')
        self.assertEqual(response.content.count(b'class="asset-grid logistics-asset-grid"'), 6)
        self.assertContains(response, 'Add a fleet asset')
        self.assertContains(response, 'Plate number or hull ID')
        self.assertContains(response, 'Choose an asset type')
        self.assertContains(response, 'Register asset')

    def test_asset_onboarding_rejects_missing_or_duplicate_details_with_friendly_message(self):
        group = Group.objects.create(name='Logistics Officers')
        user = User.objects.create_user(username='asset_logistics', password='SecurePass123!')
        user.groups.add(group)
        vehicle_type = VehicleType.objects.create(name='LAND')
        existing_vehicle = Vehicle.objects.create(
            model_name='Existing Fleet Vehicle',
            plate_number='PDR-ALS-02',
            vehicle_type=vehicle_type,
        )
        self.client.force_login(user)

        missing_fields_response = self.client.post(
            reverse('dashboard_portal:logistics_dashboard'),
            {'action': 'add_vehicle', 'model_name': 'New ambulance'},
            follow=True,
        )
        duplicate_identifier_response = self.client.post(
            reverse('dashboard_portal:logistics_dashboard'),
            {
                'action': 'add_vehicle',
                'model_name': 'Replacement ambulance',
                'plate_number': existing_vehicle.plate_number,
                'vehicle_type': vehicle_type.pk,
            },
            follow=True,
        )

        self.assertContains(
            missing_fields_response,
            'Please enter the asset name, plate number or hull ID, and asset type.',
        )
        self.assertContains(
            duplicate_identifier_response,
            "An asset with plate number or hull ID &#x27;PDR-ALS-02&#x27; is already registered.",
        )


class LogisticsReportExportTests(TestCase):
    def setUp(self):
        group = Group.objects.create(name='Logistics Officers')
        self.user = User.objects.create_user(username='report_logistics', password='SecurePass123!')
        self.user.groups.add(group)
        self.client.force_login(self.user)
        self.land_type = VehicleType.objects.create(name='LAND')
        self.vehicle = Vehicle.objects.create(
            model_name='Report Land Vehicle',
            plate_number='LAND-REPORT-01',
            vehicle_type=self.land_type,
        )
        content_type = ContentType.objects.get_for_model(Vehicle)
        self.event = LogEntry.objects.create(
            user=self.user,
            content_type=content_type,
            object_id=str(self.vehicle.pk),
            object_repr=str(self.vehicle),
            action_flag=CHANGE,
            change_message='Deployed asset unit to Puerto Princesa for Search and Rescue.',
        )
        today = timezone.localdate()
        LogEntry.objects.filter(pk=self.event.pk).update(
            action_time=timezone.make_aware(
                datetime.combine(today, datetime.min.time()),
                timezone.get_current_timezone(),
            )
        )

    def test_menu_report_form_offers_periods_and_export_formats(self):
        response = self.client.get(reverse('dashboard_portal:logistics_dashboard'))
        rendered = response.content.decode()

        self.assertContains(response, 'data-menu-action="generate-report"')
        self.assertContains(response, 'value="month">Monthly')
        self.assertContains(
            response,
            f'id="report-month" name="month" class="form-control" value="{timezone.localdate():%Y-%m}"',
        )
        self.assertContains(response, 'value="week">Weekly')
        self.assertContains(response, 'value="day">Daily')
        self.assertContains(response, 'value="custom">Custom date range')
        self.assertContains(response, 'name="report_kind"')
        self.assertContains(response, 'name="maintenance_category"')
        self.assertContains(response, 'value="damage">Damage and fault reports')
        self.assertContains(response, 'value="incident">Incident reports')
        self.assertContains(response, 'Report a fleet incident')
        self.assertContains(response, 'name="injury_status"')
        self.assertContains(response, 'name="witnesses"')
        self.assertContains(response, 'name="actions_taken"')
        self.assertContains(response, 'name="follow_up_recommendations"')
        self.assertContains(response, 'value="csv"')
        self.assertContains(response, 'value="xlsx"')
        self.assertContains(response, 'value="pdf"')
        self.assertContains(response, 'action="/fleet/reports/activity/"')
        self.assertContains(response, 'data-menu-action="file-incident"')
        self.assertContains(response, 'data-menu-action="print-incident"')
        self.assertLess(
            rendered.index('</form>', rendered.index('id="generateReportForm"')),
            rendered.index('action="/fleet/reports/incidents/"'),
        )

    def test_incident_reporting_is_available_on_every_fleet_dashboard(self):
        dashboards = (
            ('Seacraft Dispatch', 'seacraft_dispatch'),
            ('Technicians', 'repairman_dashboard'),
            ('Maritime_Tech', 'seacraft_dashboard'),
        )
        for group_name, dashboard_name in dashboards:
            with self.subTest(dashboard=dashboard_name):
                user = User.objects.create_user(username=f'{dashboard_name}_incident_user')
                user.groups.add(Group.objects.create(name=group_name))
                self.client.force_login(user)

                response = self.client.get(reverse(f'dashboard_portal:{dashboard_name}'))

                self.assertContains(response, 'Report a fleet incident')
                self.assertContains(response, 'action="/fleet/reports/incidents/"')
                self.assertContains(response, 'data-menu-action="file-incident"')
                self.assertContains(response, 'data-menu-action="print-incident"')

        self.client.force_login(self.user)

    def test_dispatch_and_mechanics_can_file_and_print_incident_reports(self):
        sea_type = VehicleType.objects.create(name='MARINE')
        role_assets = (
            ('Seacraft Dispatch', 'dispatch_incident', sea_type, 'SEA-DISPATCH-INCIDENT'),
            ('Technicians', 'land_mechanic_incident', self.land_type, 'LAND-MECHANIC-INCIDENT'),
            ('Maritime_Tech', 'sea_mechanic_incident', sea_type, 'SEA-MECHANIC-INCIDENT'),
        )
        incident_time = timezone.localtime().replace(second=0, microsecond=0)

        for group_name, username, vehicle_type, plate_number in role_assets:
            with self.subTest(role=group_name):
                user = User.objects.create_user(username=username)
                user.groups.add(Group.objects.create(name=group_name))
                vehicle = Vehicle.objects.create(
                    model_name=f'{username} asset',
                    plate_number=plate_number,
                    vehicle_type=vehicle_type,
                )
                self.client.force_login(user)

                submission = self.client.post(
                    reverse('dashboard_portal:fleet_report_incident'),
                    {
                        'vehicle': vehicle.pk,
                        'incident_type': 'DAMAGE',
                        'occurred_at': incident_time.strftime('%Y-%m-%dT%H:%M'),
                        'location': 'Incident test location',
                        'description': 'Incident filed by dashboard role.',
                        'damage_details': 'Test damage details.',
                        'injury_status': 'NO',
                        'witnesses': 'Dispatcher on duty',
                        'actions_taken': 'Asset removed from service.',
                        'follow_up_recommendations': 'Inspect before next deployment.',
                    },
                    follow=True,
                )
                self.assertContains(
                    submission,
                    'Incident report saved and added to fleet incident reports.',
                )

                printed_report = self.client.get(
                    reverse('dashboard_portal:fleet_activity_report'),
                    {
                        'report_kind': 'incident',
                        'period': 'day',
                        'day': timezone.localdate().isoformat(),
                        'format': 'print',
                    },
                )

                self.assertEqual(printed_report.status_code, 200)
                self.assertContains(printed_report, 'Incident Report')
                self.assertContains(printed_report, plate_number)
                self.assertContains(printed_report, 'window.print()')

        self.client.force_login(self.user)

    def test_report_prints_and_exports_csv_xlsx_and_pdf(self):
        today = timezone.localdate()
        iso_year, iso_week, _ = today.isocalendar()
        requests = (
            ('print', {'period': 'day', 'day': today.isoformat()}),
            ('csv', {'period': 'day', 'day': today.isoformat()}),
            ('xlsx', {'period': 'week', 'week': f'{iso_year}-W{iso_week:02d}'}),
            ('pdf', {'period': 'month', 'month': today.strftime('%Y-%m')}),
            ('csv', {
                'period': 'custom',
                'start_date': today.isoformat(),
                'end_date': today.isoformat(),
            }),
        )
        for export_format, params in requests:
            with self.subTest(export_format=export_format, period=params['period']):
                response = self.client.get(
                    reverse('dashboard_portal:logistics_generate_report'),
                    {**params, 'format': export_format},
                )
                self.assertEqual(response.status_code, 200)
                if export_format == 'print':
                    self.assertContains(response, 'Deployment Activity Report')
                    self.assertContains(response, 'window.print()')
                else:
                    self.assertIn('attachment;', response['Content-Disposition'])
                if export_format == 'csv':
                    self.assertIn('text/csv', response['Content-Type'])
                    self.assertIn(b'Report Land Vehicle', response.content)
                elif export_format == 'xlsx':
                    self.assertIn(
                        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                        response['Content-Type'],
                    )
                    from openpyxl import load_workbook

                    workbook = load_workbook(BytesIO(response.content), read_only=True)
                    self.assertEqual(workbook.active['A3'].value, today.strftime('%Y-%m-%d 00:00'))
                    self.assertIn('Report Land Vehicle', workbook.active['D3'].value)
                elif export_format == 'pdf':
                    self.assertEqual(response['Content-Type'], 'application/pdf')
                    self.assertTrue(response.content.startswith(b'%PDF'))

    def test_report_rejects_invalid_range_and_unsupported_format(self):
        reverse_url = reverse('dashboard_portal:logistics_generate_report')
        invalid_range = self.client.get(
            reverse_url,
            {'format': 'csv', 'period': 'custom', 'start_date': '2026-10-09', 'end_date': '2026-10-01'},
        )
        unsupported_format = self.client.get(
            reverse_url,
            {'format': 'doc', 'period': 'day', 'day': timezone.localdate().isoformat()},
        )

        self.assertEqual(invalid_range.status_code, 400)
        self.assertContains(unsupported_format, 'Choose Print, CSV, Excel, or PDF', status_code=400)

    def test_report_requires_logistics_access(self):
        self.client.force_login(User.objects.create_user(username='unauthorized_report_user'))

        response = self.client.get(
            reverse('dashboard_portal:logistics_generate_report'),
            {'format': 'csv', 'period': 'day', 'day': timezone.localdate().isoformat()},
        )

        self.assertRedirects(response, reverse('dashboard_portal:homepage'))

    def test_seacraft_dispatch_can_generate_seacraft_activity_report(self):
        sea_group = Group.objects.create(name='Seacraft Dispatch')
        dispatch_user = User.objects.create_user(username='dispatch_report_user')
        dispatch_user.groups.add(sea_group)
        sea_type = VehicleType.objects.create(name='MARINE')
        sea_vehicle = Vehicle.objects.create(
            model_name='Report Seacraft',
            plate_number='SEA-REPORT-01',
            vehicle_type=sea_type,
        )
        LogEntry.objects.create(
            user=self.user,
            content_type=ContentType.objects.get_for_model(Vehicle),
            object_id=str(sea_vehicle.pk),
            object_repr=str(sea_vehicle),
            action_flag=CHANGE,
            change_message='Deployed asset unit to Honda Bay for Search and Rescue.',
        )
        self.client.force_login(dispatch_user)

        response = self.client.get(
            reverse('dashboard_portal:fleet_activity_report'),
            {'format': 'csv', 'period': 'day', 'day': timezone.localdate().isoformat()},
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn(b'Report Seacraft', response.content)
        self.assertNotIn(b'Report Land Vehicle', response.content)

    def test_logistics_can_submit_and_export_fleet_incident_report(self):
        assigned_driver = Driver.objects.create(
            name='Last Assigned Land Driver',
            license_number='LAND-LAST-001',
        )
        self.vehicle.assigned_driver = assigned_driver
        self.vehicle.save(update_fields=['assigned_driver'])
        sea_vehicle = Vehicle.objects.create(
            model_name='Report Seacraft Incident',
            plate_number='SEA-INCIDENT-01',
            vehicle_type=VehicleType.objects.create(name='MARINE'),
        )
        LogEntry.objects.create(
            user=self.user,
            content_type=ContentType.objects.get_for_model(Vehicle),
            object_id=str(sea_vehicle.pk),
            object_repr=str(sea_vehicle),
            action_flag=CHANGE,
            change_message=(
                'Deployed asset unit to Honda Bay for Search and Rescue '
                'with operator Last Assigned Sea Operator.'
            ),
        )
        incident_time = timezone.localtime().replace(second=0, microsecond=0)
        for vehicle, location, description in (
            (self.vehicle, 'Puerto Princesa', 'Vehicle struck debris during response.'),
            (sea_vehicle, 'Honda Bay', 'Vessel equipment was damaged during response.'),
        ):
            response = self.client.post(
                reverse('dashboard_portal:fleet_report_incident'),
                {
                    'vehicle': vehicle.pk,
                    'incident_type': 'DAMAGE',
                    'occurred_at': incident_time.strftime('%Y-%m-%dT%H:%M'),
                    'location': location,
                    'description': description,
                    'damage_details': 'Equipment damage recorded.',
                    'injury_status': 'YES',
                    'injury_details': 'Operator assessed by the medical response team.',
                    'witnesses': 'Response team lead',
                    'actions_taken': 'Asset secured and supervisor notified.',
                    'follow_up_recommendations': 'Complete maintenance inspection.',
                },
                follow=True,
            )
            self.assertRedirects(response, reverse('dashboard_portal:logistics_dashboard'))
            self.assertContains(response, 'Incident report saved and added to fleet incident reports.')

        incidents = {
            incident.vehicle_id: incident
            for incident in FleetIncident.objects.select_related('vehicle')
        }
        self.assertEqual(len(incidents), 2)
        self.assertEqual(incidents[self.vehicle.pk].reported_by, self.user)
        self.assertEqual(incidents[self.vehicle.pk].last_assigned_driver, 'Last Assigned Land Driver')
        self.assertEqual(incidents[sea_vehicle.pk].last_assigned_driver, 'Last Assigned Sea Operator')
        self.assertEqual(
            incidents[self.vehicle.pk].injury_details,
            'Operator assessed by the medical response team.',
        )
        self.assertEqual(incidents[self.vehicle.pk].witnesses, 'Response team lead')
        self.assertEqual(
            incidents[self.vehicle.pk].actions_taken,
            'Asset secured and supervisor notified.',
        )
        self.assertEqual(
            incidents[self.vehicle.pk].follow_up_recommendations,
            'Complete maintenance inspection.',
        )

        report_url = reverse('dashboard_portal:fleet_activity_report')
        report_params = {
            'report_kind': 'incident',
            'period': 'day',
            'day': timezone.localdate().isoformat(),
        }
        for report_format in ('print', 'csv', 'xlsx', 'pdf'):
            with self.subTest(format=report_format):
                report = self.client.get(
                    report_url,
                    {**report_params, 'format': report_format},
                )
                self.assertEqual(report.status_code, 200)
                if report_format == 'print':
                    self.assertContains(report, 'Incident Report')
                    self.assertContains(report, 'Puerto Princesa')
                    self.assertContains(report, 'Last Assigned Land Driver')
                    self.assertContains(report, 'Last Assigned Sea Operator')
                    self.assertContains(report, 'Provincial Disaster Risk Reduction and Management Office')
                    self.assertContains(report, 'Prepared by / Date')
                    self.assertContains(report, 'Print Incident Report')
                    self.assertContains(report, 'Witnesses')
                    self.assertContains(report, 'Response team lead')
                    self.assertContains(report, 'Operator assessed by the medical response team.')
                    self.assertContains(report, 'Asset secured and supervisor notified.')
                    self.assertContains(report, 'Complete maintenance inspection.')
                elif report_format == 'csv':
                    self.assertIn(b'Incident Report', report.content)
                    self.assertIn(b'Puerto Princesa', report.content)
                    self.assertIn(b'Last Assigned Land Driver', report.content)
                    self.assertIn(b'Last Assigned Sea Operator', report.content)
                    self.assertIn(
                        b'Injury / medical details: Operator assessed by the medical response team.',
                        report.content,
                    )
                    self.assertIn(b'Witnesses: Response team lead', report.content)
                    self.assertIn(b'Immediate actions: Asset secured and supervisor notified.', report.content)
                    self.assertIn(b'Follow-up recommendations: Complete maintenance inspection.', report.content)
                elif report_format == 'xlsx':
                    self.assertIn(
                        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                        report['Content-Type'],
                    )
                    self.assertIn('fleet-incident-', report['Content-Disposition'])
                    from openpyxl import load_workbook

                    workbook = load_workbook(BytesIO(report.content), read_only=True)
                    report_details = (
                        workbook.active['E3'].value,
                        workbook.active['E4'].value,
                    )
                    self.assertTrue(any('Puerto Princesa' in value for value in report_details))
                    self.assertEqual(workbook.active['G3'].value, 'Last Assigned Sea Operator')
                    self.assertEqual(workbook.active['G4'].value, 'Last Assigned Land Driver')
                    self.assertIn('Witnesses: Response team lead', workbook.active['E4'].value)
                    self.assertIn(
                        'Operator assessed by the medical response team.',
                        workbook.active['E4'].value,
                    )
                    self.assertIn(
                        'Immediate actions: Asset secured and supervisor notified.',
                        workbook.active['E4'].value,
                    )
                else:
                    self.assertEqual(report['Content-Type'], 'application/pdf')
                    self.assertTrue(report.content.startswith(b'%PDF'))
                    self.assertIn('fleet-incident-', report['Content-Disposition'])

    def test_incident_form_requires_details_when_injuries_are_reported(self):
        response = self.client.post(
            reverse('dashboard_portal:fleet_report_incident'),
            {
                'vehicle': self.vehicle.pk,
                'incident_type': 'ACCIDENT',
                'occurred_at': timezone.localtime().strftime('%Y-%m-%dT%H:%M'),
                'location': 'Test location',
                'description': 'Incident requiring injury details.',
                'injury_status': 'YES',
                'injury_details': '',
            },
            follow=True,
        )

        self.assertContains(
            response,
            'Describe the injuries or medical assistance required.',
        )
        self.assertFalse(FleetIncident.objects.exists())

    def test_damage_report_includes_logged_mechanic_faults(self):
        LogEntry.objects.create(
            user=self.user,
            content_type=ContentType.objects.get_for_model(Vehicle),
            object_id=str(self.vehicle.pk),
            object_repr=str(self.vehicle),
            action_flag=CHANGE,
            change_message='Reported maintenance fault for Report Land Vehicle: Body / onboard equipment: Door damaged.',
        )

        response = self.client.get(
            reverse('dashboard_portal:fleet_activity_report'),
            {
                'report_kind': 'damage',
                'period': 'day',
                'day': timezone.localdate().isoformat(),
                'format': 'csv',
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn(b'Damage / fault reported', response.content)
        self.assertIn(b'Door damaged', response.content)

    def test_seacraft_technician_cannot_report_incident_for_land_asset(self):
        maritime_group = Group.objects.create(name='Maritime_Tech')
        maritime_user = User.objects.create_user(username='incident_sea_tech')
        maritime_user.groups.add(maritime_group)
        self.client.force_login(maritime_user)

        response = self.client.post(
            reverse('dashboard_portal:fleet_report_incident'),
            {
                'vehicle': self.vehicle.pk,
                'incident_type': 'DAMAGE',
                'occurred_at': timezone.localtime().strftime('%Y-%m-%dT%H:%M'),
                'location': 'Test location',
                'description': 'Unauthorized land asset incident.',
                'damage_details': '',
                'injury_status': 'NO',
            },
            follow=True,
        )

        self.assertFalse(FleetIncident.objects.exists())
        self.assertContains(response, 'Fleet asset: Select a valid choice')

    def test_unauthorized_user_cannot_submit_incident_report(self):
        self.client.force_login(User.objects.create_user(username='blocked_incident_reporter'))

        response = self.client.post(
            reverse('dashboard_portal:fleet_report_incident'),
            {
                'vehicle': self.vehicle.pk,
                'incident_type': 'DAMAGE',
                'occurred_at': timezone.localtime().strftime('%Y-%m-%dT%H:%M'),
                'location': 'Test location',
                'description': 'Unauthorized incident.',
                'injury_status': 'NO',
            },
        )

        self.assertRedirects(response, reverse('dashboard_portal:homepage'))
        self.assertFalse(FleetIncident.objects.exists())


class DispatchDashboardAlignmentTests(TestCase):
    def setUp(self):
        self.marine_type = VehicleType.objects.create(name='MARINE')
        self.maritime_type = VehicleType.objects.create(name='MARITIME')
        self.land_type = VehicleType.objects.create(name='LAND')
        self.seacraft_user = User.objects.create_user(
            username='maritime_dispatcher',
            password='SecurePass123!',
        )
        logistics_group = Group.objects.create(name='Logistics Officers')
        self.logistics_user = User.objects.create_user(
            username='fleet_logistics',
            password='SecurePass123!',
        )
        self.logistics_user.groups.add(logistics_group)
        self.operator = Driver.objects.create(
            name='Qualified Maritime Operator',
            license_number='MAR-001',
        )

    def create_vehicle(self, *, model_name, plate_number, vehicle_type, **kwargs):
        return Vehicle.objects.create(
            model_name=model_name,
            plate_number=plate_number,
            vehicle_type=vehicle_type,
            **kwargs,
        )

    def test_logistics_can_add_land_drivers_and_seacraft_operators(self):
        self.client.force_login(self.logistics_user)
        for name, operator_type, license_number, expected_authority in (
            ('New Land Driver', 'LAND', 'N01-23-456789', 'LTO'),
            ('New Sea Operator', 'SEA', 'MRA-100', 'MARINA'),
        ):
            with self.subTest(operator=name):
                response = self.client.post(
                    reverse('dashboard_portal:logistics_dashboard'),
                    {
                        'action': 'add_operator',
                        'operator_type': operator_type,
                        'name': name,
                        'license_number': license_number,
                        'phone_number': '555-0100',
                    },
                )
                self.assertRedirects(response, reverse('dashboard_portal:logistics_dashboard'))
                operator = Driver.objects.get(name=name)
                self.assertEqual(operator.license_authority, expected_authority)
                self.assertEqual(operator.license_number, license_number)
                self.assertEqual(operator.phone_number, '555-0100')

    def test_logistics_sets_authority_from_operator_type_without_user_input(self):
        self.client.force_login(self.logistics_user)
        for operator_type, name, authority in (
            ('SEA', 'Automatically Maritime Authority', 'MARINA'),
            ('LAND', 'Automatically Land Authority', 'LTO'),
        ):
            with self.subTest(operator_type=operator_type):
                response = self.client.post(
                    reverse('dashboard_portal:logistics_dashboard'),
                    {
                        'action': 'add_operator',
                        'operator_type': operator_type,
                        'name': name,
                        'license_number': 'TEST-123456',
                        'license_authority': 'LTO' if operator_type == 'SEA' else 'MARINA',
                    },
                )

                self.assertRedirects(response, reverse('dashboard_portal:logistics_dashboard'))
                operator = Driver.objects.get(name=name)
                self.assertEqual(operator.license_authority, authority)

    def test_logistics_preserves_license_number_as_entered(self):
        self.client.force_login(self.logistics_user)
        response = self.client.post(
            reverse('dashboard_portal:logistics_dashboard'),
            {
                'action': 'add_operator',
                'operator_type': 'SEA',
                'name': 'Marina License Operator',
                'license_number': 'MRA-601',
            },
        )

        self.assertRedirects(response, reverse('dashboard_portal:logistics_dashboard'))
        operator = Driver.objects.get(name='Marina License Operator')
        self.assertEqual(operator.license_authority, 'MARINA')
        self.assertEqual(operator.license_number, 'MRA-601')

    def test_logistics_assignment_drawer_saves_on_dropdown_change(self):
        self.create_vehicle(
            model_name='Assignment Drawer Test Vehicle',
            plate_number='LAND-DRAWER-01',
            vehicle_type=self.land_type,
        )
        self.client.force_login(self.logistics_user)
        response = self.client.get(reverse('dashboard_portal:logistics_dashboard'))

        self.assertNotContains(response, 'Save All Operator Assignments')
        self.assertContains(response, 'onchange="this.form.submit()"')
        self.assertContains(response, 'name="action" value="set_driver"')

    def test_logistics_dropdown_changes_save_land_and_seacraft_assignments(self):
        land_operator = Driver.objects.create(
            name='Bulk Land Operator',
            license_number='LAND-500',
        )
        land_vehicle = self.create_vehicle(
            model_name='Bulk Assigned Land Vehicle',
            plate_number='LAND-BULK-01',
            vehicle_type=self.land_type,
        )
        seacraft = self.create_vehicle(
            model_name='Bulk Assigned Seacraft',
            plate_number='SEA-BULK-01',
            vehicle_type=self.marine_type,
        )
        self.client.force_login(self.logistics_user)

        land_response = self.client.post(
            reverse('dashboard_portal:logistics_dashboard'),
            {
                'action': 'set_driver',
                'vehicle_id': land_vehicle.pk,
                'driver_id': land_operator.pk,
            },
        )
        sea_response = self.client.post(
            reverse('dashboard_portal:logistics_dashboard'),
            {
                'action': 'set_driver',
                'vehicle_id': seacraft.pk,
                'driver_id': self.operator.pk,
            },
        )

        self.assertRedirects(land_response, reverse('dashboard_portal:logistics_dashboard'))
        self.assertRedirects(sea_response, reverse('dashboard_portal:logistics_dashboard'))
        land_vehicle.refresh_from_db()
        seacraft.refresh_from_db()
        self.assertEqual(land_vehicle.assigned_driver, land_operator)
        self.assertEqual(seacraft.assigned_driver, self.operator)

    def test_logistics_dropdown_rejects_operator_already_assigned_elsewhere(self):
        first_vehicle = self.create_vehicle(
            model_name='First Bulk Land Vehicle',
            plate_number='LAND-BULK-02',
            vehicle_type=self.land_type,
        )
        second_vehicle = self.create_vehicle(
            model_name='Second Bulk Land Vehicle',
            plate_number='LAND-BULK-03',
            vehicle_type=self.land_type,
        )
        land_operator = Driver.objects.create(
            name='Duplicate Bulk Land Operator',
            license_number='LAND-501',
        )
        self.client.force_login(self.logistics_user)

        response = self.client.post(
            reverse('dashboard_portal:logistics_dashboard'),
            {
                'action': 'set_driver',
                'vehicle_id': first_vehicle.pk,
                'driver_id': land_operator.pk,
            },
        )
        self.assertRedirects(response, reverse('dashboard_portal:logistics_dashboard'))

        response = self.client.post(
            reverse('dashboard_portal:logistics_dashboard'),
            {
                'action': 'set_driver',
                'vehicle_id': second_vehicle.pk,
                'driver_id': land_operator.pk,
            },
            follow=True,
        )

        self.assertContains(response, 'already assigned to active fleet asset')
        first_vehicle.refresh_from_db()
        second_vehicle.refresh_from_db()
        self.assertEqual(first_vehicle.assigned_driver, land_operator)
        self.assertIsNone(second_vehicle.assigned_driver)

    def test_seacraft_dispatch_can_add_only_maritime_operators(self):
        self.client.force_login(self.seacraft_user)
        response = self.client.post(
            reverse('dashboard_portal:seacraft_dispatch'),
            {
                'action': 'add_operator',
                'name': 'Dispatch Added Operator',
                'license_number': 'MRA-300',
                'phone_number': '555-0300',
            },
        )

        self.assertRedirects(response, reverse('dashboard_portal:seacraft_dispatch'))
        operator = Driver.objects.get(name='Dispatch Added Operator')
        self.assertEqual(operator.license_authority, 'MARINA')
        self.assertEqual(operator.license_number, 'MRA-300')

    def test_seacraft_dispatch_assigns_marina_authority_automatically(self):
        self.client.force_login(self.seacraft_user)
        response = self.client.post(
            reverse('dashboard_portal:seacraft_dispatch'),
            {
                'action': 'add_operator',
                'name': 'Automatically Maritime Authority',
                'license_number': 'MRA-301',
                'license_authority': 'LTO',
            },
        )

        self.assertRedirects(response, reverse('dashboard_portal:seacraft_dispatch'))
        operator = Driver.objects.get(name='Automatically Maritime Authority')
        self.assertEqual(operator.license_authority, 'MARINA')

    def test_maritime_asset_classification_matches_across_dashboards(self):
        marine_craft = self.create_vehicle(
            model_name='Marine Name Craft',
            plate_number='SEA-MARINE-01',
            vehicle_type=self.marine_type,
        )
        maritime_craft = self.create_vehicle(
            model_name='Maritime Name Craft',
            plate_number='SEA-MARITIME-01',
            vehicle_type=self.maritime_type,
        )
        land_vehicle = self.create_vehicle(
            model_name='Land Name Vehicle',
            plate_number='LAND-01',
            vehicle_type=self.land_type,
        )

        self.client.force_login(self.logistics_user)
        logistics_response = self.client.get(reverse('dashboard_portal:logistics_dashboard'))
        self.assertEqual(logistics_response.status_code, 200)
        self.assertEqual(
            {craft.id for craft in logistics_response.context['sea_crafts']},
            {marine_craft.id, maritime_craft.id},
        )
        self.assertEqual(
            {vehicle.id for vehicle in logistics_response.context['land_vehicles']},
            {land_vehicle.id},
        )

        self.client.force_login(self.seacraft_user)
        dispatch_response = self.client.get(reverse('dashboard_portal:seacraft_dispatch'))
        self.assertEqual(dispatch_response.status_code, 200)
        for craft in (marine_craft, maritime_craft):
            self.assertContains(dispatch_response, craft.model_name)
        self.assertNotContains(dispatch_response, land_vehicle.model_name)

    def test_public_homepage_links_to_deployment_activity_log(self):
        response = self.client.get(reverse('dashboard_portal:homepage'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, reverse('dashboard_portal:deployment_activity_log'))
        self.assertContains(response, 'Deployment Activity Log')

    def test_public_deployment_activity_log_shows_deployment_and_return(self):
        craft = self.create_vehicle(
            model_name='Deployment History Craft',
            plate_number='SEA-HISTORY-01',
            vehicle_type=self.marine_type,
        )
        self.client.force_login(self.seacraft_user)
        deploy_response = self.client.post(
            reverse('dashboard_portal:seacraft_dispatch'),
            {
                'action': 'deploy_vehicle',
                'vehicle_id': craft.pk,
                'driver_id': self.operator.pk,
                'deployment_location': 'Honda Bay',
                'deployment_purpose': 'Search and Rescue',
                'deployment_time': '2030-01-02T10:30',
            },
        )
        self.assertRedirects(deploy_response, reverse('dashboard_portal:seacraft_dispatch'))

        return_response = self.client.post(
            reverse('dashboard_portal:seacraft_dispatch'),
            {'action': 'return_vehicle', 'vehicle_id': craft.pk},
        )
        self.assertRedirects(return_response, reverse('dashboard_portal:seacraft_dispatch'))
        self.client.logout()

        response = self.client.get(reverse('dashboard_portal:deployment_activity_log'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Deployment Activity Log')
        self.assertContains(response, 'Deployment History Craft')
        self.assertContains(response, 'Honda Bay')
        self.assertContains(response, 'Search and Rescue')
        self.assertContains(response, 'RETURNED')
        self.assertContains(response, 'fa-ship text-info')
        self.assertContains(response, 'fa-arrow-up-right-from-square text-primary')
        self.assertContains(response, 'fa-arrow-rotate-left text-success')
        self.assertEqual(response.context['deployment_event_count'], 2)

    def test_deployment_log_refresh_returns_partial_updates_without_page_reload(self):
        land_vehicle = self.create_vehicle(
            model_name='Deployment Refresh Land Asset',
            plate_number='LAND-REFRESH-01',
            vehicle_type=self.land_type,
        )
        land_driver = Driver.objects.create(
            name='Deployment Refresh Driver',
            license_number='LTO-123',
        )
        self.client.force_login(self.logistics_user)
        self.client.post(
            reverse('dashboard_portal:logistics_dashboard'),
            {
                'action': 'deploy_vehicle',
                'vehicle_id': land_vehicle.pk,
                'driver_id': land_driver.pk,
                'deployment_location': 'Puerto Princesa',
                'deployment_purpose': 'Logistics Support',
                'deployment_time': '2030-01-02T10:30',
            },
        )

        response = self.client.get(
            reverse('dashboard_portal:deployment_activity_log'),
            HTTP_X_REQUESTED_WITH='XMLHttpRequest',
        )

        self.assertEqual(response.status_code, 200)
        payload = json.loads(response.content)
        self.assertEqual(payload['count'], 1)
        self.assertIn('Deployment Refresh Land Asset', payload['rows'])
        self.assertIn('fa-truck text-warning', payload['rows'])
        self.assertIn('fa-arrow-up-right-from-square text-primary', payload['rows'])
        page_response = self.client.get(reverse('dashboard_portal:deployment_activity_log'))
        self.assertContains(page_response, 'setInterval(refreshDeploymentActivity, 30000)')
        self.assertContains(page_response, 'rows.dataset.signature !== data.signature')
        self.assertContains(page_response, 'if (refreshInProgress || document.hidden) return')

    def test_seacraft_dispatch_uses_logistics_status_tabs(self):
        operational = self.create_vehicle(
            model_name='Tabbed Operational Craft',
            plate_number='SEA-TAB-01',
            vehicle_type=self.marine_type,
        )
        deployed = self.create_vehicle(
            model_name='Tabbed Deployed Craft',
            plate_number='SEA-TAB-02',
            vehicle_type=self.marine_type,
            status='DEPLOYED',
        )
        maintenance = self.create_vehicle(
            model_name='Tabbed Maintenance Craft',
            plate_number='SEA-TAB-03',
            vehicle_type=self.marine_type,
            status='MAINTENANCE',
        )
        pending_disposal = self.create_vehicle(
            model_name='Tabbed Disposal Craft',
            plate_number='SEA-TAB-04',
            vehicle_type=self.marine_type,
            status='PENDING_DISPOSAL',
        )

        self.client.force_login(self.seacraft_user)
        response = self.client.get(reverse('dashboard_portal:seacraft_dispatch'))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['sea_standby'], [operational])
        self.assertEqual(response.context['sea_deployed'], [deployed])
        self.assertEqual(
            response.context['sea_maintenance'],
            [maintenance, pending_disposal],
        )
        self.assertContains(response, 'Standby &amp; Operational (1)')
        self.assertContains(response, 'Active Deployments (1)')
        self.assertContains(response, 'Maintenance &amp; Disposal (2)')
        self.assertEqual(response.content.count(b'data-horizontal-card-scroll="true"'), 3)

    def test_seacraft_dispatch_header_actions_are_in_menu_drawer(self):
        self.client.force_login(self.seacraft_user)

        response = self.client.get(reverse('dashboard_portal:seacraft_dispatch'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'id="seacraftMenuDrawer"')
        self.assertContains(response, 'data-menu-action="print"')
        self.assertContains(response, 'data-menu-action="add-operator"')
        self.assertContains(response, 'Connected to Dispatch Server')
        self.assertContains(response, 'Log Out')
        self.assertContains(response, 'aria-label="Open seacraft dispatch menu"')
        self.assertContains(response, 'data-menu-action="generate-report"')
        self.assertContains(response, 'Generate Fleet Report')

    def test_unauthorized_seacraft_dispatch_user_redirects_to_homepage(self):
        unauthorized_user = User.objects.create_user(username='unauthorized_dispatch')
        self.client.force_login(unauthorized_user)

        response = self.client.get(reverse('dashboard_portal:seacraft_dispatch'))

        self.assertRedirects(response, reverse('dashboard_portal:homepage'))

    def test_dispatch_deployment_details_are_visible_in_logistics_dashboard(self):
        craft = self.create_vehicle(
            model_name='Shared Dispatch Craft',
            plate_number='SEA-SHARED-01',
            vehicle_type=self.maritime_type,
        )
        self.client.force_login(self.seacraft_user)

        response = self.client.post(
            reverse('dashboard_portal:seacraft_dispatch'),
            {
                'action': 'deploy_vehicle',
                'vehicle_id': craft.id,
                'driver_id': self.operator.id,
                'deployment_location': 'Honda Bay',
                'deployment_purpose': 'Emergency Search & Rescue',
                'deployment_time': '2030-01-02T10:30',
            },
        )

        self.assertRedirects(response, reverse('dashboard_portal:seacraft_dispatch'))
        craft.refresh_from_db()
        self.assertEqual(craft.status, 'DEPLOYED')
        self.assertEqual(craft.assigned_driver, self.operator)
        self.assertEqual(craft.deployment_location, 'Honda Bay')
        self.assertEqual(craft.deployment_purpose, 'Emergency Search & Rescue')
        self.assertIsNotNone(craft.deployment_time)

        self.client.force_login(self.logistics_user)
        dashboard = self.client.get(reverse('dashboard_portal:logistics_dashboard'))
        self.assertContains(dashboard, 'Shared Dispatch Craft')
        self.assertContains(dashboard, 'Honda Bay')
        self.assertContains(dashboard, 'Emergency Search &amp; Rescue')
        self.assertContains(dashboard, 'Qualified Maritime Operator')

    def test_logistics_deployment_details_are_visible_in_seacraft_dispatch(self):
        craft = self.create_vehicle(
            model_name='Logistics Deployed Craft',
            plate_number='SEA-LOGISTICS-01',
            vehicle_type=self.marine_type,
        )
        self.client.force_login(self.logistics_user)

        response = self.client.post(
            reverse('dashboard_portal:logistics_dashboard'),
            {
                'action': 'deploy_vehicle',
                'vehicle_id': craft.id,
                'driver_id': self.operator.id,
                'deployment_location': 'Coron Bay',
                'deployment_purpose': 'Coastal & Sea Patrol / Reconnaissance',
                'deployment_time': '2030-01-03T08:15',
            },
        )

        self.assertRedirects(response, reverse('dashboard_portal:logistics_dashboard'))
        craft.refresh_from_db()
        self.assertEqual(craft.status, 'DEPLOYED')
        self.assertEqual(craft.assigned_driver, self.operator)
        self.assertEqual(craft.deployment_location, 'Coron Bay')
        self.assertEqual(craft.deployment_purpose, 'Coastal & Sea Patrol / Reconnaissance')
        self.assertIsNotNone(craft.deployment_time)

        self.client.force_login(self.seacraft_user)
        dispatch_dashboard = self.client.get(reverse('dashboard_portal:seacraft_dispatch'))
        self.assertContains(dispatch_dashboard, 'Logistics Deployed Craft')
        self.assertContains(dispatch_dashboard, 'Coron Bay')
        self.assertContains(dispatch_dashboard, 'Coastal &amp; Sea Patrol / Reconnaissance')

    def test_dispatch_rejects_nonoperational_craft_and_unqualified_operator(self):
        craft = self.create_vehicle(
            model_name='Blocked Dispatch Craft',
            plate_number='SEA-BLOCKED-01',
            vehicle_type=self.marine_type,
            status='MAINTENANCE',
        )
        land_operator = Driver.objects.create(
            name='Land Driver',
            license_number='LAND-001',
        )
        self.client.force_login(self.seacraft_user)

        response = self.client.post(
            reverse('dashboard_portal:seacraft_dispatch'),
            {
                'action': 'deploy_vehicle',
                'vehicle_id': craft.id,
                'driver_id': self.operator.id,
                'deployment_location': 'Honda Bay',
                'deployment_purpose': 'Emergency Search & Rescue',
                'deployment_time': '2030-01-02T10:30',
            },
            follow=True,
        )
        self.assertContains(response, 'must be operational')
        craft.refresh_from_db()
        self.assertEqual(craft.status, 'MAINTENANCE')

        craft.status = 'OPERATIONAL'
        craft.save(update_fields=['status'])
        response = self.client.post(
            reverse('dashboard_portal:seacraft_dispatch'),
            {
                'action': 'deploy_vehicle',
                'vehicle_id': craft.id,
                'driver_id': land_operator.id,
                'deployment_location': 'Honda Bay',
                'deployment_purpose': 'Emergency Search & Rescue',
                'deployment_time': '2030-01-02T10:30',
            },
            follow=True,
        )
        self.assertContains(response, 'qualified seacraft operator')
        craft.refresh_from_db()
        self.assertEqual(craft.status, 'OPERATIONAL')

    def test_operator_can_be_reassigned_from_maintenance_and_return_clears_deployment(self):
        repair_craft = self.create_vehicle(
            model_name='Maintenance Craft',
            plate_number='SEA-MAINT-01',
            vehicle_type=self.marine_type,
            status='MAINTENANCE',
            assigned_driver=self.operator,
        )
        operational_craft = self.create_vehicle(
            model_name='Ready Craft',
            plate_number='SEA-READY-01',
            vehicle_type=self.maritime_type,
        )
        self.client.force_login(self.seacraft_user)

        response = self.client.post(
            reverse('dashboard_portal:seacraft_dispatch'),
            {
                'action': 'deploy_vehicle',
                'vehicle_id': operational_craft.id,
                'driver_id': self.operator.id,
                'deployment_location': 'Honda Bay',
                'deployment_purpose': 'Emergency Search & Rescue',
                'deployment_time': '2030-01-02T10:30',
            },
        )
        self.assertRedirects(response, reverse('dashboard_portal:seacraft_dispatch'))
        repair_craft.refresh_from_db()
        operational_craft.refresh_from_db()
        self.assertIsNone(repair_craft.assigned_driver)
        self.assertEqual(operational_craft.assigned_driver, self.operator)

        response = self.client.post(
            reverse('dashboard_portal:seacraft_dispatch'),
            {'action': 'return_vehicle', 'vehicle_id': operational_craft.id},
        )
        self.assertRedirects(response, reverse('dashboard_portal:seacraft_dispatch'))
        operational_craft.refresh_from_db()
        self.assertEqual(operational_craft.status, 'OPERATIONAL')
        self.assertIsNone(operational_craft.deployment_location)
        self.assertIsNone(operational_craft.deployment_purpose)
        self.assertIsNone(operational_craft.deployment_time)


class MechanicMaintenanceWorkflowTests(TestCase):
    def setUp(self):
        self.land_type = VehicleType.objects.create(name='LAND')
        self.marine_type = VehicleType.objects.create(name='MARINE')
        technician_group = Group.objects.create(name='Technicians')
        maritime_group = Group.objects.create(name='Maritime_Tech')
        self.land_mechanic = User.objects.create_user(
            username='land_mechanic',
            password='SecurePass123!',
        )
        self.land_mechanic.groups.add(technician_group)
        self.sea_mechanic = User.objects.create_user(
            username='sea_mechanic',
            password='SecurePass123!',
        )
        self.sea_mechanic.groups.add(maritime_group)
        self.land_vehicle = Vehicle.objects.create(
            model_name='Land Work Vehicle',
            plate_number='LAND-WORK-01',
            vehicle_type=self.land_type,
        )
        self.seacraft = Vehicle.objects.create(
            model_name='Sea Work Vessel',
            plate_number='SEA-WORK-01',
            vehicle_type=self.marine_type,
        )

    def post_as_mechanic(self, vehicle, action, *, is_seacraft=False, **data):
        user = self.sea_mechanic if is_seacraft else self.land_mechanic
        url_name = 'dashboard_portal:seacraft_dashboard' if is_seacraft else 'dashboard_portal:repairman_dashboard'
        if action == 'REPORT_FAULT' and 'fault_description' in data:
            data[f'fault-{vehicle.pk}-fault_description'] = data.pop('fault_description')
            data[f'fault-{vehicle.pk}-fault_type'] = data.pop('fault_type', 'engine')
        if action == 'COMPLETE_MAINTENANCE' and 'maintenance_checklist' in data:
            data[f'checklist-{vehicle.pk}-maintenance_checklist'] = data.pop('maintenance_checklist')
        if action == 'SCHEDULE_MAINTENANCE':
            data[f'schedule-{vehicle.pk}-maintenance_type'] = data.pop('maintenance_type', 'TIRES')
            data[f'schedule-{vehicle.pk}-due_date'] = data.pop('due_date')
            if 'description' in data:
                data[f'schedule-{vehicle.pk}-description'] = data.pop('description')
        self.client.force_login(user)
        return self.client.post(
            reverse(url_name),
            {'vehicle_id': vehicle.id, 'action_type': action, **data},
        )

    def test_mechanic_dashboards_offer_report_generation(self):
        for user, dashboard in (
            (self.land_mechanic, 'dashboard_portal:repairman_dashboard'),
            (self.sea_mechanic, 'dashboard_portal:seacraft_dashboard'),
        ):
            with self.subTest(dashboard=dashboard):
                self.client.force_login(user)
                response = self.client.get(reverse(dashboard))

                self.assertEqual(response.status_code, 200)
                self.assertContains(response, 'id="mechanicMenuDrawer"')
                self.assertContains(response, 'aria-label="Open mechanic menu"')
                self.assertContains(response, 'Maintenance Menu')
                self.assertContains(response, 'Maintenance Workbench')
                self.assertContains(response, 'data-menu-action="print"')
                self.assertContains(response, 'data-menu-action="generate-report"')
                self.assertContains(response, 'Log Out')
                self.assertContains(response, 'Generate Fleet Report')
                self.assertContains(response, 'value="custom">Custom date range')
                self.assertContains(response, 'value="xlsx"')
                self.assertContains(response, 'Regular maintenance (tires and oil)')

    def test_unauthorized_mechanic_users_redirect_to_homepage(self):
        unauthorized_user = User.objects.create_user(username='blocked_user')
        self.client.force_login(unauthorized_user)

        for dashboard in (
            'dashboard_portal:repairman_dashboard',
            'dashboard_portal:seacraft_dashboard',
        ):
            with self.subTest(dashboard=dashboard):
                response = self.client.get(reverse(dashboard))
                self.assertRedirects(response, reverse('dashboard_portal:homepage'))

    def test_mechanic_report_is_limited_to_its_asset_division(self):
        content_type = ContentType.objects.get_for_model(Vehicle)
        LogEntry.objects.create(
            user=self.land_mechanic,
            content_type=content_type,
            object_id=str(self.land_vehicle.pk),
            object_repr=str(self.land_vehicle),
            action_flag=CHANGE,
            change_message='Deployed asset unit to Puerto Princesa for Search and Rescue.',
        )
        LogEntry.objects.create(
            user=self.sea_mechanic,
            content_type=content_type,
            object_id=str(self.seacraft.pk),
            object_repr=str(self.seacraft),
            action_flag=CHANGE,
            change_message='Deployed asset unit to Honda Bay for Search and Rescue.',
        )
        report_url = reverse('dashboard_portal:fleet_activity_report')
        params = {
            'format': 'csv',
            'period': 'day',
            'day': timezone.localdate().isoformat(),
        }

        for user, included_asset, excluded_asset in (
            (self.land_mechanic, b'Land Work Vehicle', b'Sea Work Vessel'),
            (self.sea_mechanic, b'Sea Work Vessel', b'Land Work Vehicle'),
        ):
            with self.subTest(username=user.username):
                self.client.force_login(user)
                response = self.client.get(report_url, params)

                self.assertEqual(response.status_code, 200)
                self.assertIn(included_asset, response.content)
                self.assertNotIn(excluded_asset, response.content)

    def test_maintenance_reports_include_regular_and_other_work_types(self):
        content_type = ContentType.objects.get_for_model(Vehicle)
        events = (
            (self.land_vehicle, 'Scheduled Tire replacement for 2026-10-09: Replace front tires.'),
            (self.land_vehicle, 'Scheduled Oil replacement for 2026-10-09: Replace engine oil.'),
            (self.land_vehicle, 'Scheduled Other maintenance for 2026-10-09: Repair body panel.'),
            (self.land_vehicle, 'Reported maintenance fault for Land Work Vehicle: Engine: Belt noise'),
            (self.land_vehicle, 'Completed maintenance checklist and returned Land Work Vehicle to operational status.'),
            (self.seacraft, 'Scheduled Tire replacement for 2026-10-09: Replace trailer tires.'),
        )
        for vehicle, message in events:
            LogEntry.objects.create(
                user=self.land_mechanic,
                content_type=content_type,
                object_id=str(vehicle.pk),
                object_repr=str(vehicle),
                action_flag=CHANGE,
                change_message=message,
            )
        report_url = reverse('dashboard_portal:fleet_activity_report')
        base_params = {
            'report_kind': 'maintenance',
            'format': 'csv',
            'period': 'day',
            'day': timezone.localdate().isoformat(),
        }

        self.client.force_login(self.land_mechanic)
        regular_report = self.client.get(
            report_url,
            {**base_params, 'maintenance_category': 'regular'},
        )
        other_report = self.client.get(
            report_url,
            {**base_params, 'maintenance_category': 'other'},
        )
        all_report = self.client.get(
            report_url,
            {**base_params, 'maintenance_category': 'all'},
        )

        self.assertEqual(regular_report.status_code, 200)
        self.assertIn(b'Regular maintenance', regular_report.content)
        self.assertIn(b'Tire replacement', regular_report.content)
        self.assertIn(b'Oil replacement', regular_report.content)
        self.assertNotIn(b'body panel', regular_report.content)
        self.assertNotIn(b'Sea Work Vessel', regular_report.content)
        self.assertEqual(other_report.status_code, 200)
        self.assertIn(b'body panel', other_report.content)
        self.assertIn(b'Belt noise', other_report.content)
        self.assertIn(b'Completed maintenance checklist', other_report.content)
        self.assertNotIn(b'Tire replacement', other_report.content)
        self.assertEqual(all_report.status_code, 200)
        self.assertIn(b'Tire replacement', all_report.content)
        self.assertIn(b'body panel', all_report.content)

    def test_maintenance_report_supports_print_and_file_exports(self):
        LogEntry.objects.create(
            user=self.land_mechanic,
            content_type=ContentType.objects.get_for_model(Vehicle),
            object_id=str(self.land_vehicle.pk),
            object_repr=str(self.land_vehicle),
            action_flag=CHANGE,
            change_message='Reported maintenance fault for Land Work Vehicle: Engine: Belt noise',
        )
        report_url = reverse('dashboard_portal:fleet_activity_report')
        params = {
            'report_kind': 'maintenance',
            'maintenance_category': 'other',
            'period': 'day',
            'day': timezone.localdate().isoformat(),
        }

        self.client.force_login(self.land_mechanic)
        for export_format in ('print', 'csv', 'xlsx', 'pdf'):
            with self.subTest(format=export_format):
                response = self.client.get(
                    report_url,
                    {**params, 'format': export_format},
                )
                self.assertEqual(response.status_code, 200)
                if export_format == 'print':
                    self.assertContains(response, 'Other maintenance and repairs Report')
                    self.assertContains(response, 'Belt noise')
                elif export_format == 'pdf':
                    self.assertEqual(response['Content-Type'], 'application/pdf')
                    self.assertTrue(response.content.startswith(b'%PDF'))
                    self.assertIn('fleet-maintenance-', response['Content-Disposition'])
                elif export_format == 'xlsx':
                    self.assertIn(
                        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                        response['Content-Type'],
                    )
                    self.assertIn('fleet-maintenance-', response['Content-Disposition'])
                else:
                    self.assertIn(b'Other maintenance and repairs', response.content)
                    self.assertIn(b'Belt noise', response.content)

    def test_fault_description_is_required_before_maintenance(self):
        for vehicle, is_seacraft in (
            (self.land_vehicle, False),
            (self.seacraft, True),
        ):
            with self.subTest(asset=vehicle.model_name):
                self.post_as_mechanic(vehicle, 'REPORT_FAULT', is_seacraft=is_seacraft, fault_description='   ')
                vehicle.refresh_from_db()
                self.assertEqual(vehicle.status, 'OPERATIONAL')
                self.assertEqual(vehicle.maintenance_problem, '')

    def test_fault_report_moves_asset_to_maintenance_and_persists_fault(self):
        for vehicle, is_seacraft in (
            (self.land_vehicle, False),
            (self.seacraft, True),
        ):
            with self.subTest(asset=vehicle.model_name):
                self.post_as_mechanic(
                    vehicle,
                    'REPORT_FAULT',
                    is_seacraft=is_seacraft,
                    fault_description='Engine makes a loud grinding noise.',
                )
                vehicle.refresh_from_db()
                self.assertEqual(vehicle.status, 'MAINTENANCE')
                self.assertEqual(
                    vehicle.maintenance_problem,
                    'Engine: Engine makes a loud grinding noise.',
                )

    def test_fault_report_requires_a_valid_issue_type(self):
        self.post_as_mechanic(
            self.land_vehicle,
            'REPORT_FAULT',
            fault_type='unsupported',
            fault_description='A fault was reported.',
        )

        self.land_vehicle.refresh_from_db()
        self.assertEqual(self.land_vehicle.status, 'OPERATIONAL')
        self.assertEqual(self.land_vehicle.maintenance_problem, '')

    def test_mechanic_can_schedule_future_preventive_maintenance(self):
        due_date = timezone.localdate() + timedelta(days=30)
        response = self.post_as_mechanic(
            self.land_vehicle,
            'SCHEDULE_MAINTENANCE',
            maintenance_type='OIL',
            due_date=due_date.isoformat(),
            description='Use the approved engine oil.',
        )

        self.assertRedirects(response, reverse('dashboard_portal:repairman_dashboard'))
        self.land_vehicle.refresh_from_db()
        self.assertEqual(self.land_vehicle.status, 'OPERATIONAL')
        self.assertEqual(self.land_vehicle.scheduled_maintenance_type, 'OIL')
        self.assertEqual(self.land_vehicle.scheduled_maintenance_date, due_date)
        self.assertEqual(
            self.land_vehicle.scheduled_maintenance_description,
            'Use the approved engine oil.',
        )

    def test_overdue_scheduled_maintenance_moves_vehicle_to_maintenance(self):
        driver = Driver.objects.create(name='Scheduled Service Driver')
        self.land_vehicle.status = 'DEPLOYED'
        self.land_vehicle.assigned_driver = driver
        self.land_vehicle.deployment_location = 'Test location'
        self.land_vehicle.deployment_purpose = 'Test deployment'
        self.land_vehicle.scheduled_maintenance_type = 'TIRES'
        self.land_vehicle.scheduled_maintenance_date = timezone.localdate() - timedelta(days=1)
        self.land_vehicle.scheduled_maintenance_description = 'Replace front tires.'
        self.land_vehicle.save()
        self.client.force_login(self.land_mechanic)

        response = self.client.get(reverse('dashboard_portal:repairman_dashboard'))

        self.assertEqual(response.status_code, 200)
        self.land_vehicle.refresh_from_db()
        self.assertEqual(self.land_vehicle.status, 'MAINTENANCE')
        self.assertEqual(
            self.land_vehicle.maintenance_problem,
            'Scheduled Tire replacement is due. Replace front tires.',
        )
        self.assertIsNone(self.land_vehicle.assigned_driver)
        self.assertIsNone(self.land_vehicle.deployment_location)
        self.assertIsNone(self.land_vehicle.scheduled_maintenance_date)

    def test_overdue_vehicle_cannot_be_deployed_from_logistics(self):
        logistics_group = Group.objects.create(name='Logistics Officers')
        logistics_user = User.objects.create_user(
            username='fleet_logistics',
            password='SecurePass123!',
        )
        logistics_user.groups.add(logistics_group)
        self.land_vehicle.scheduled_maintenance_type = 'OIL'
        self.land_vehicle.scheduled_maintenance_date = timezone.localdate()
        self.land_vehicle.save()
        self.client.force_login(logistics_user)

        response = self.client.post(
            reverse('dashboard_portal:logistics_dashboard'),
            {
                'action': 'deploy_vehicle',
                'vehicle_id': self.land_vehicle.id,
            },
        )

        self.assertRedirects(response, reverse('dashboard_portal:logistics_dashboard'))
        self.land_vehicle.refresh_from_db()
        self.assertEqual(self.land_vehicle.status, 'MAINTENANCE')

    def test_past_maintenance_schedule_date_is_rejected(self):
        self.post_as_mechanic(
            self.land_vehicle,
            'SCHEDULE_MAINTENANCE',
            maintenance_type='OIL',
            due_date=(timezone.localdate() - timedelta(days=1)).isoformat(),
        )

        self.land_vehicle.refresh_from_db()
        self.assertEqual(self.land_vehicle.status, 'OPERATIONAL')
        self.assertIsNone(self.land_vehicle.scheduled_maintenance_date)

    def test_seacraft_mechanic_can_schedule_and_trigger_due_maintenance(self):
        due_date = timezone.localdate() + timedelta(days=1)
        self.post_as_mechanic(
            self.seacraft,
            'SCHEDULE_MAINTENANCE',
            is_seacraft=True,
            maintenance_type='OTHER',
            due_date=due_date.isoformat(),
            description='Inspect bilge pump.',
        )

        self.seacraft.refresh_from_db()
        self.assertEqual(self.seacraft.status, 'OPERATIONAL')
        self.assertEqual(self.seacraft.scheduled_maintenance_type, 'OTHER')
        self.assertEqual(self.seacraft.scheduled_maintenance_date, due_date)

        self.seacraft.scheduled_maintenance_date = timezone.localdate()
        self.seacraft.save(update_fields=['scheduled_maintenance_date'])
        self.client.force_login(self.sea_mechanic)
        response = self.client.get(reverse('dashboard_portal:seacraft_dashboard'))

        self.assertEqual(response.status_code, 200)
        self.seacraft.refresh_from_db()
        self.assertEqual(self.seacraft.status, 'MAINTENANCE')
        self.assertEqual(
            self.seacraft.maintenance_problem,
            'Scheduled Other maintenance is due. Inspect bilge pump.',
        )

    def test_deployed_assets_show_mission_location_and_click_to_call_driver(self):
        for vehicle, is_seacraft, phone_number in (
            (self.land_vehicle, False, '+639171234567'),
            (self.seacraft, True, '+639189876543'),
        ):
            with self.subTest(asset=vehicle.model_name):
                driver = Driver.objects.create(
                    name=f'{vehicle.model_name} Operator',
                    phone_number=phone_number,
                    license_number='MAR-100' if is_seacraft else 'LAND-100',
                )
                vehicle.status = 'DEPLOYED'
                vehicle.assigned_driver = driver
                vehicle.deployment_location = 'Honda Bay'
                vehicle.deployment_purpose = 'Emergency Search & Rescue'
                vehicle.save()

                user = self.sea_mechanic if is_seacraft else self.land_mechanic
                dashboard = (
                    'dashboard_portal:seacraft_dashboard'
                    if is_seacraft else 'dashboard_portal:repairman_dashboard'
                )
                self.client.force_login(user)
                response = self.client.get(reverse(dashboard))

                self.assertEqual(response.status_code, 200)
                self.assertContains(response, 'Location:')
                self.assertContains(response, 'Honda Bay')
                self.assertContains(response, 'Mission detail:')
                self.assertContains(response, 'Emergency Search &amp; Rescue')
                self.assertContains(response, f'Driver: {driver.name}')
                self.assertContains(response, f'href="tel:{phone_number}"')
                self.assertContains(response, phone_number)

    def test_reporting_a_deployed_fault_clears_dispatch_assignment(self):
        driver = Driver.objects.create(name='Assigned Land Operator')
        self.land_vehicle.status = 'DEPLOYED'
        self.land_vehicle.assigned_driver = driver
        self.land_vehicle.deployment_location = 'Test location'
        self.land_vehicle.deployment_purpose = 'Test deployment'
        self.land_vehicle.save()

        self.post_as_mechanic(
            self.land_vehicle,
            'REPORT_FAULT',
            fault_description='The brakes make a grinding noise.',
        )

        self.land_vehicle.refresh_from_db()
        self.assertEqual(self.land_vehicle.status, 'MAINTENANCE')
        self.assertIsNone(self.land_vehicle.assigned_driver)
        self.assertIsNone(self.land_vehicle.deployment_location)
        self.assertIsNone(self.land_vehicle.deployment_purpose)

    def test_every_asset_specific_checklist_item_is_required(self):
        for vehicle, is_seacraft in (
            (self.land_vehicle, False),
            (self.seacraft, True),
        ):
            with self.subTest(asset=vehicle.model_name):
                vehicle.status = 'MAINTENANCE'
                vehicle.maintenance_problem = 'User reported a fault.'
                vehicle.save()
                self.post_as_mechanic(
                    vehicle,
                    'COMPLETE_MAINTENANCE',
                    is_seacraft=is_seacraft,
                    maintenance_checklist=['repairs_complete'],
                )
                vehicle.refresh_from_db()
                self.assertEqual(vehicle.status, 'MAINTENANCE')
                self.assertEqual(vehicle.maintenance_problem, 'User reported a fault.')

    def test_complete_checklist_returns_land_and_sea_assets_to_service(self):
        for vehicle, is_seacraft in (
            (self.land_vehicle, False),
            (self.seacraft, True),
        ):
            with self.subTest(asset=vehicle.model_name):
                vehicle.status = 'MAINTENANCE'
                vehicle.maintenance_problem = 'User reported a fault.'
                vehicle.save()
                self.post_as_mechanic(
                    vehicle,
                    'COMPLETE_MAINTENANCE',
                    is_seacraft=is_seacraft,
                    maintenance_checklist=[
                        'repairs_complete',
                        'safety_systems_checked',
                        'test_run_passed',
                        'no_unresolved_issues',
                        'propulsion_fuel_hull_checked' if is_seacraft else 'fluids_battery_tires_checked',
                    ],
                )
                vehicle.refresh_from_db()
                self.assertEqual(vehicle.status, 'OPERATIONAL')
                self.assertEqual(vehicle.maintenance_problem, '')

    def test_legacy_status_toggle_cannot_bypass_mechanic_workflow(self):
        for vehicle, is_seacraft in (
            (self.land_vehicle, False),
            (self.seacraft, True),
        ):
            with self.subTest(asset=vehicle.model_name):
                self.post_as_mechanic(
                    vehicle,
                    'STATUS_TOGGLE',
                    is_seacraft=is_seacraft,
                    status='MAINTENANCE',
                )
                vehicle.refresh_from_db()
                self.assertEqual(vehicle.status, 'OPERATIONAL')

    def test_seacraft_dispatch_cannot_bypass_maintenance_checklist(self):
        self.seacraft.status = 'MAINTENANCE'
        self.seacraft.maintenance_problem = 'User reported a fault.'
        self.seacraft.save()
        self.client.force_login(self.sea_mechanic)

        self.client.post(
            reverse('dashboard_portal:seacraft_dispatch'),
            {
                'vehicle_id': self.seacraft.id,
                'action': 'STATUS_TOGGLE',
                'status': 'OPERATIONAL',
            },
        )

        self.seacraft.refresh_from_db()
        self.assertEqual(self.seacraft.status, 'MAINTENANCE')
        self.assertEqual(self.seacraft.maintenance_problem, 'User reported a fault.')

    def test_mechanic_pages_show_fault_prompt_and_no_status_toggles(self):
        for user, vehicle, url_name, is_seacraft in (
            (self.land_mechanic, self.land_vehicle, 'dashboard_portal:repairman_dashboard', False),
            (self.sea_mechanic, self.seacraft, 'dashboard_portal:seacraft_dashboard', True),
        ):
            with self.subTest(asset=vehicle.model_name):
                self.client.force_login(user)
                response = self.client.get(reverse(url_name))
                self.assertContains(response, 'Description of issue')
                if not is_seacraft:
                    self.assertContains(response, 'Operational Fleet')
                    self.assertContains(response, 'Maintenance')
                    self.assertContains(response, 'Issue type')
                    self.assertContains(response, 'Driver:')
                    self.assertContains(response, 'data-land-fleet-tab="maintenance"')
                    self.assertContains(response, 'data-land-fleet-group="operational" hidden')
                else:
                    self.assertContains(response, 'data-sea-fleet-tab="maintenance"')
                    self.assertContains(response, 'data-sea-fleet-group="operational" hidden')
                    self.assertContains(response, 'Schedule other maintenance')
                    self.assertContains(response, 'Issue type')
                    self.assertContains(response, 'Driver: Unassigned')
                if not is_seacraft:
                    self.assertContains(response, f'data-bs-target="#landMaintenanceModal{vehicle.id}"')
                    self.assertContains(response, 'modal-dialog modal-dialog-centered modal-lg')
                else:
                    self.assertContains(response, f'data-bs-target="#seaMaintenanceModal{vehicle.id}"')
                    self.assertContains(response, 'modal-dialog modal-dialog-centered modal-lg')
                vehicle.status = 'MAINTENANCE'
                vehicle.maintenance_problem = 'User-reported test fault.'
                vehicle.save()
                response = self.client.get(reverse(url_name))
                self.assertContains(response, 'Return-to-service checklist')
                self.assertContains(response, 'User-reported test fault.')
                if not is_seacraft:
                    self.assertContains(response, 'data-land-fleet-group="maintenance">')
                else:
                    self.assertContains(response, 'data-sea-fleet-group="maintenance">')
                self.assertContains(response, 'class="maintenance-checklist"')
                self.assertNotContains(response, 'STATUS_TOGGLE')
